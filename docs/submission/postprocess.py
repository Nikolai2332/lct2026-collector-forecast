"""Постобработка .docx после pandoc: рамки таблиц, повтор шапки таблицы на каждой странице, шрифт 9 пт в таблицах,
ширина таблиц — по ширине страницы, разрыв страницы после титула (перед оглавлением)."""

import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

path = sys.argv[1]
d = Document(path)


def set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")) + tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "9A9A9A")
        b.append(e)
    tblPr.append(b)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), "5000")
    w.set(qn("w:type"), "pct")
    tblPr.append(w)


for t in d.tables:
    set_borders(t)
    rows = t.rows
    if rows:
        trPr = rows[0]._tr.get_or_add_trPr()
        h = OxmlElement("w:tblHeader")
        h.set(qn("w:val"), "true")
        trPr.append(h)
        for cell in rows[0].cells:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "E8EEF6")
            cell._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        # строку таблицы не разрывать между страницами (иначе на новой странице — пустой «хвост» строки)
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cs)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.size = Pt(9)

# Разрыв страницы перед оглавлением (после титульного блока)
for el in d.element.body.iter(qn("w:sdt")):
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    el.addprevious(p)
    break
d.save(path)
print(f"postprocess: таблиц {len(d.tables)}")
