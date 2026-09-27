"""Проверка собранной документации (в контейнере collector-docs): .docx открывается, считает заголовки, таблицы и
рисунки, ищет пустые заголовки и запрещённое в тексте (.docx и текст .pdf).

    python3 check_docx.py <файл.docx> <текст pdf.txt>
Код 1 — если найдено запрещённое или пустые заголовки."""

import re
import sys

from docx import Document

docx_path, pdf_txt = sys.argv[1], sys.argv[2]
d = Document(docx_path)
heads = [p for p in d.paragraphs if p.style.name.lower().startswith("heading")]
empty = [p for p in heads if not p.text.strip()]
pics = len(d.inline_shapes)
print(f"заголовков: {len(heads)} (пустых {len(empty)}), таблиц: {len(d.tables)}, рисунков: {pics}")
for level in ("heading 1", "heading 2", "heading 3"):
    print(f"  {level}: {sum(1 for p in heads if p.style.name.lower() == level)}")

text = "\n".join(p.text for p in d.paragraphs)
text += "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
text += open(pdf_txt, encoding="utf-8").read()
FORBIDDEN = [
    (r"C:\\Users|C:/Users", "путь к личной папке"),
    (r"collector_test", "тестовая база"),
    (r"сертифицирован|гарантир", "«сертифицировано/гарантирует»"),
    (r"PROMPT_|промпт", "рабочие файлы и промпты"),
    (r"STATUS\.md|tz_general|tz_backend|tz_frontend|tz_ml", "рабочие документы"),
    (r"backups/|(?<![\w/])data/", "рабочие каталоги"),
    (r"(?<![\w.])\.env(?![.\w])", "файл секретов .env"),
    (r"сесси[яиюей]|этап(е|а|ом)? №", "описание хода работы"),
    (r"change-me|dev-secret|JWT_SECRET=\S+|POSTGRES_PASSWORD=\S+", "секреты"),
    (r"\[\[|\]\]|<!--", "служебная разметка"),
]
bad = 0
for pattern, what in FORBIDDEN:
    hits = sorted({m.group(0) for m in re.finditer(pattern, text, flags=re.I)})
    if hits:
        bad += 1
        print(f"НАЙДЕНО ({what}): {hits[:5]}")
print("запрещённого не найдено" if not bad else f"проблем: {bad}")
sys.exit(1 if bad or empty else 0)
