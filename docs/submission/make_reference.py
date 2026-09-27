"""Шаблон стилей .docx для pandoc: шрифты PT Serif / PT Sans (кириллица), размеры, отступы, заголовок 1 — с новой
страницы, колонтитул с названием и номером страницы, поля A4. Исходник — стандартный reference.docx pandoc."""

import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

src, dst = sys.argv[1], sys.argv[2]
d = Document(src)
BODY, HEAD = "PT Serif", "PT Sans"


def font(style, name, size=None, bold=None, color=None):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if color:
        f.color.rgb = RGBColor.from_string(color)


_all = {s.name.lower(): s for s in d.styles}


class _Styles:
    """Стили по имени без учёта регистра: в reference.docx pandoc они «heading 1», а не «Heading 1»."""

    def __getitem__(self, name):
        return _all[name.lower()]


styles = _Styles()
names = {n for n in _all} | {s.name for s in d.styles}
for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
    if name.lower() in _all:
        font(styles[name], BODY, 11)
        pf = styles[name].paragraph_format
        pf.space_after = Pt(4)
        pf.line_spacing = 1.1
for name, size in (("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12), ("Title", 26), ("Subtitle", 15)):
    font(styles[name], HEAD, size, True, "1F3A5F")
styles["Heading 1"].paragraph_format.page_break_before = True
styles["Heading 1"].paragraph_format.space_after = Pt(10)
for name in ("Heading 2", "Heading 3"):
    styles[name].paragraph_format.space_before = Pt(12)
    styles[name].paragraph_format.keep_with_next = True
for name in ("Author", "Date"):
    if name.lower() in _all:
        font(styles[name], HEAD, 13)
        styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
styles["Title"].paragraph_format.space_before = Pt(150)
for name in ("Image Caption", "Table Caption", "Caption"):
    if name.lower() in _all:
        font(styles[name], HEAD, 9.5, False, "404040")
        styles[name].font.italic = False
for name in ("Source Code", "Verbatim Char"):
    if name.lower() in _all:
        font(styles[name], "PT Mono", 9)

sec = d.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
sec.different_first_page_header_footer = True  # на титуле колонтитула нет


def field(par, instr):
    """Поле Word (PAGE): begin — instrText — separate — значение — end."""
    def run_with(el):
        r = par.add_run()
        r._r.append(el)
        return r

    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    run_with(fc)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run_with(it)
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    run_with(fc)
    par.add_run("1")
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    run_with(fc)


hp = sec.header.paragraphs[0]
hp.text = "Сервис прогноза отказов датчиков в подземных коллекторах · ЛЦТ-2026, «Отказ датчика»"
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs:
    r.font.size, r.font.name = Pt(8.5), HEAD
    r.font.color.rgb = RGBColor.from_string("808080")
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run("Страница ")
field(fp, "PAGE")
for r in fp.runs:
    r.font.size, r.font.name = Pt(9), HEAD
d.save(dst)
print("reference:", dst)
