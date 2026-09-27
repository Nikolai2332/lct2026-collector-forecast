"""Ручной импорт CSV/XLSX и чистка событий по правилам из ТЗ.

Для гигабайтных журналов этот путь не годится — там будет ETL через DuckDB и COPY (app/etl).
"""

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time

from openpyxl import load_workbook
from pydantic import BaseModel, Field, ValidationError, field_validator
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import xmlio
from app.dbcompat import upsert
from app.labels import risk_level_from_health
from app.models import Channel, EventRecent, Object, Prediction
from app.timeutil import to_msk_naive

MAX_ERRORS = 50
# Больше строк за раз — через ETL (app/etl): ручной импорт держит всё в памяти одной транзакцией
MAX_IMPORT_ROWS = 300_000
# XLSX — это zip: сжатый файл в 50 МБ может распаковаться в гигабайты («zip-бомба»)
MAX_XLSX_UNPACKED_BYTES = 300 * 1024 * 1024
MAX_XLSX_ENTRIES = 2000

# Статусы, которые ML считает отказом (цель модели)
FAULT_STATUSES = {"Неисправен", "Отключено устройство"}

# Заголовки файлов заказчика (после нормализации) → поля
HEADER_ALIASES = {
    "objects": {
        "ид_объект": "id", "id": "id", "object_id": "id",
        "иерархия_уровень": "level", "level": "level",
        "родитель": "parent_id", "parent_id": "parent_id", "parent": "parent_id",
        "вид_объекта": "kind", "kind": "kind",
        "название": "name", "name": "name",
    },
    "channels": {
        "ид_канала_данных": "id", "id": "id", "channel_id": "id",
        "тип_инж_системы": "system_type", "тип_инженерной_системы": "system_type", "system_type": "system_type",
        "тип_датчика": "sensor_type", "sensor_type": "sensor_type",
        "тег_инженерной_системы": "tag", "тег": "tag", "tag": "tag",
        "название_датчика": "name", "name": "name",
        "ид_объект": "object_id", "object_id": "object_id",
    },
    "events": {
        "ид_события": "event_id", "event_id": "event_id",
        "ид_канала_данных": "channel_id", "channel_id": "channel_id",
        "дата": "date", "date": "date",
        "время": "time", "time": "time",
        "тревожное": "is_alarm", "is_alarm": "is_alarm",
        "значение_датчика": "value", "value": "value",
        "ts": "ts", "timestamp": "ts",
    },
    "predictions": {
        k: k for k in ("channel_id", "at", "horizon_h", "prob", "health", "risk_level", "top_factors", "model_version")
    },
}
REQUIRED_FIELDS = {
    "objects": {"id", "level", "kind", "name"},
    "channels": {"id", "system_type", "sensor_type", "tag", "name", "object_id"},
    "events": {"channel_id", "is_alarm", "value"},
    "predictions": {"channel_id", "at", "prob", "model_version"},
}

_DATE_VALUE = re.compile(r"^\d{2}\.\d{2}\.\d{4}(\s+\d{1,2}:\d{2}(:\d{2})?)?$")
_TRUE = {"t", "true", "1", "да", "yes", "y"}
_FALSE = {"f", "false", "0", "нет", "no", "n", ""}


@dataclass
class ImportReport:
    kind: str
    file_name: str
    rows_total: int = 0
    rows_imported: int = 0
    rows_skipped: int = 0
    errors: list[dict] = field(default_factory=list)
    imported_predictions: list[Prediction] = field(default_factory=list)

    def error(self, row: int, message: str) -> None:
        self.rows_skipped += 1
        if len(self.errors) < MAX_ERRORS:
            self.errors.append({"row": row, "message": message})


# ---------- Чтение файлов ----------


def normalize_header(name: str) -> str:
    # Заголовки из PDF/Excel приходят с пробелами вокруг подчёркиваний и BOM
    name = str(name or "").replace("﻿", "").strip().lower()
    name = re.sub(r"[\s\-]+", "_", name)
    return re.sub(r"_+", "_", name).strip("_")


def safe_file_name(name: str | None) -> str:
    """Имя файла для отчёта и журнала: без пути, управляющих символов и не длиннее 128 символов."""
    base = re.split(r"[\\/]", str(name or ""))[-1]
    base = re.sub(r"[\x00-\x1f\x7f]", "", base).strip()
    return base[:128] or "upload"


def _check_xlsx(content: bytes) -> None:
    if not content.startswith(b"PK\x03\x04"):
        raise ValueError("Файл .xlsx повреждён или это не Excel-книга")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile:
        raise ValueError("Файл .xlsx повреждён или это не Excel-книга")
    if len(infos) > MAX_XLSX_ENTRIES or sum(i.file_size for i in infos) > MAX_XLSX_UNPACKED_BYTES:
        raise ValueError("Книга Excel слишком большая после распаковки — загрузите данные через ETL")


def _check_rows(n: int) -> None:
    if n > MAX_IMPORT_ROWS:
        raise ValueError(f"В файле больше {MAX_IMPORT_ROWS:,} строк — загрузите его через ETL".replace(",", " "))


def read_table(content: bytes, file_name: str) -> list[dict[str, str]]:
    """CSV (UTF-8, разделитель , или ;), XLSX или XML → список словарей со строковыми значениями.

    Тип проверяется и по расширению, и по содержимому; число строк и размер распакованной книги ограничены."""
    lower = file_name.lower()
    if lower.endswith(".xlsx"):
        _check_xlsx(content)
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            rows = ws.iter_rows(values_only=True)
            header = [normalize_header(h) for h in next(rows, [])]
            result = []
            for r in rows:
                if r is None or all(v is None for v in r):
                    continue
                result.append({header[i]: _cell_to_str(v) for i, v in enumerate(r) if i < len(header)})
                _check_rows(len(result))
        finally:
            wb.close()
        return result
    if lower.endswith(".csv"):
        if content.startswith(b"PK\x03\x04") or b"\x00" in content[:8192]:
            raise ValueError("Файл .csv на самом деле не текстовый (похоже на Excel или архив)")
        text = content.decode("utf-8-sig")
        sample = text[:4096]
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        # Значение датчика читаем строкой всегда: в файле за 2022 год оно без кавычек
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        header = [normalize_header(h) for h in next(reader, [])]
        result = []
        for r in reader:
            if any(cell.strip() for cell in r):
                result.append({header[i]: v for i, v in enumerate(r) if i < len(header)})
                _check_rows(len(result))
        return result
    if lower.endswith(".xml"):
        if not content.lstrip(b"\xef\xbb\xbf \t\r\n").startswith(b"<"):
            raise ValueError("Файл .xml на самом деле не XML")
        result = []
        for rec in xmlio.records(xmlio.parse(content)):
            if any(str(v).strip() for v in rec.values()):
                result.append({normalize_header(k): v for k, v in rec.items()})
                _check_rows(len(result))
        return result
    raise ValueError("Поддерживаются только файлы .csv, .xlsx и .xml")


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def map_columns(rows: list[dict], kind: str) -> list[dict]:
    aliases = HEADER_ALIASES[kind]
    if not rows:
        return []
    present = {aliases[h] for h in rows[0] if h in aliases}
    missing = REQUIRED_FIELDS[kind] - present
    if kind == "events" and "ts" not in present and not {"date", "time"} <= present:
        missing |= {"date", "time"}
    if missing:
        raise ValueError(f"В файле нет обязательных колонок: {', '.join(sorted(missing))}")
    return [{aliases[k]: v for k, v in r.items() if k in aliases} for r in rows]


# ---------- Разбор значений ----------


def parse_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    s = str(value).strip().lower()
    if s in _TRUE:
        return True
    if s in _FALSE:
        return False
    raise ValueError(f"Непонятное значение флага «тревожное»: {value!r}")


def parse_date(value: str) -> date:
    s = str(value).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Непонятная дата: {value!r}")


def parse_time(value: str) -> time:
    s = str(value).strip()
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            pass
    raise ValueError(f"Непонятное время: {value!r}")


def parse_datetime(value) -> datetime:
    if isinstance(value, datetime):
        return to_msk_naive(value)
    s = str(value).strip()
    try:
        return to_msk_naive(datetime.fromisoformat(s))
    except ValueError:
        pass
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    raise ValueError(f"Непонятная дата и время: {value!r}")


@dataclass
class CleanEvent:
    source_event_id: int | None
    channel_id: int
    ts: datetime
    is_alarm: bool
    value_raw: str
    value_num: float | None
    value_text: str | None


def clean_event(channel_id, ts: datetime, is_alarm, value, event_id=None) -> CleanEvent | None:
    """Правила чистки из ТЗ. None — значение-дата, такое событие выбрасываем."""
    raw = str(value).strip() if value is not None else ""
    if not raw or _DATE_VALUE.match(raw):
        return None
    num: float | None
    try:
        num = float(raw.replace(",", "."))
        text = None
    except ValueError:
        num, text = None, raw
    return CleanEvent(
        source_event_id=int(event_id) if event_id not in (None, "") else None,
        channel_id=int(channel_id),
        ts=ts.replace(microsecond=0),
        is_alarm=parse_bool(is_alarm),
        value_raw=raw[:128],
        value_num=num,
        value_text=text[:128] if text else None,
    )


def store_events(db: Session, events: list[CleanEvent]) -> tuple[int, int, list[int]]:
    """Сортирует по времени, убирает дубли (канал, время, значение) в пачке и в БД.

    Возвращает (вставлено, дублей, затронутые каналы)."""
    events.sort(key=lambda e: (e.ts, e.channel_id))
    unique: dict[tuple, CleanEvent] = {}
    for e in events:
        unique.setdefault((e.channel_id, e.ts, e.value_raw), e)
    rows = [e.__dict__ for e in unique.values()]
    inserted = upsert(db, EventRecent.__table__, rows, ["channel_id", "ts", "value_raw"])
    return inserted, len(events) - inserted, sorted({e.channel_id for e in unique.values()})


def known_channel_ids(db: Session, ids: set[int]) -> set[int]:
    known: set[int] = set()
    ids_list = list(ids)
    for i in range(0, len(ids_list), 5000):
        known.update(db.scalars(select(Channel.id).where(Channel.id.in_(ids_list[i : i + 5000]))))
    return known


# ---------- Импорт по видам ----------


def split_tag(tag: str) -> list[str | None]:
    parts = [p for p in str(tag).strip().strip(".").split(".") if p]
    # «847-1.1.131.2.» → первый уровень до дефиса отдельно: 847, 1, 1, 131, 2
    if parts and "-" in parts[0]:
        head, _, rest = parts[0].partition("-")
        parts = [head, rest, *parts[1:]]
    parts = parts[:5]
    return parts + [None] * (5 - len(parts))


def _import_objects(db: Session, rows: list[dict], report: ImportReport) -> None:
    clean = []
    for i, r in enumerate(rows, start=2):
        try:
            clean.append({
                "id": int(r["id"]),
                "level": int(r["level"]),
                "parent_id": int(r["parent_id"]) if str(r.get("parent_id", "")).strip() not in ("", "0", "None") else None,
                "kind": r["kind"].strip(),
                "name": r["name"].strip(),
            })
        except (KeyError, ValueError, AttributeError) as e:
            report.error(i, f"Некорректная строка объекта: {e}")
    clean.sort(key=lambda o: o["level"])  # родители раньше детей — ради внешнего ключа
    report.rows_imported = upsert(db, Object.__table__, clean, ["id"], ["level", "parent_id", "kind", "name"])


def _import_channels(db: Session, rows: list[dict], report: ImportReport) -> None:
    object_ids = set(db.scalars(select(Object.id)))
    clean = []
    for i, r in enumerate(rows, start=2):
        try:
            obj = int(r["object_id"])
            if obj not in object_ids:
                report.error(i, f"Объект {obj} не найден в справочнике объектов")
                continue
            levels = split_tag(r["tag"])
            clean.append({
                "id": int(r["id"]),
                "object_id": obj,
                "system_type": r["system_type"].strip(),
                "sensor_type": r["sensor_type"].strip(),
                "tag": r["tag"].strip(),
                "name": r["name"].strip(),
                **{f"tag_l{n + 1}": levels[n] for n in range(5)},
            })
        except (KeyError, ValueError, AttributeError) as e:
            report.error(i, f"Некорректная строка датчика: {e}")
    cols = ["object_id", "system_type", "sensor_type", "tag", "name", *[f"tag_l{n}" for n in range(1, 6)]]
    report.rows_imported = upsert(db, Channel.__table__, clean, ["id"], cols)


def _import_events(db: Session, rows: list[dict], report: ImportReport) -> None:
    parsed: list[tuple[int, CleanEvent]] = []
    for i, r in enumerate(rows, start=2):
        try:
            ts = parse_datetime(r["ts"]) if r.get("ts") else datetime.combine(parse_date(r["date"]), parse_time(r["time"]))
            ev = clean_event(r["channel_id"], ts, r["is_alarm"], r.get("value"), r.get("event_id"))
        except (KeyError, ValueError) as e:
            report.error(i, str(e))
            continue
        if ev is None:
            report.error(i, "Вместо показания записана дата — строка пропущена")
            continue
        parsed.append((i, ev))
    known = known_channel_ids(db, {e.channel_id for _, e in parsed})
    valid = []
    for i, e in parsed:
        if e.channel_id in known:
            valid.append(e)
        else:
            report.error(i, f"Канал {e.channel_id} не найден в справочнике")
    inserted, duplicates, _ = store_events(db, valid)
    report.rows_imported = inserted
    report.rows_skipped += duplicates


class PredictionRow(BaseModel):
    """Контракт с ML: channel_id, at, horizon_h (всегда 24), prob, health, risk_level, top_factors, model_version."""

    channel_id: int
    at: datetime
    horizon_h: int = Field(24, gt=0)
    prob: float = Field(ge=0, le=1)
    health: int | None = Field(None, ge=0, le=100)
    risk_level: str | None = None
    top_factors: list = []
    model_version: str = Field(min_length=1, max_length=64)

    @field_validator("top_factors", mode="before")
    @classmethod
    def _parse_json(cls, v):
        if v in (None, ""):
            return []
        if isinstance(v, str):
            v = json.loads(v)
        return v if isinstance(v, list) else [v]

    @field_validator("health", "risk_level", "horizon_h", mode="before")
    @classmethod
    def _empty_to_none(cls, v):
        return None if v == "" else v

    def to_row(self) -> dict:
        health = self.health if self.health is not None else round(100 * (1 - self.prob))
        risk = self.risk_level or risk_level_from_health(health)
        if risk not in ("normal", "attention", "risk", "critical"):
            raise ValueError(f"Неизвестный уровень риска: {risk}")
        return {
            "channel_id": self.channel_id,
            "at": to_msk_naive(self.at),
            "horizon_h": self.horizon_h or 24,
            "prob": self.prob,
            "health": health,
            "risk_level": risk,
            "top_factors": self.top_factors,
            "model_version": self.model_version,
        }


def _import_predictions(db: Session, rows: list[dict], report: ImportReport) -> None:
    clean: list[tuple[int, dict]] = []
    for i, r in enumerate(rows, start=2):
        try:
            clean.append((i, PredictionRow.model_validate(r).to_row()))
        except (ValidationError, ValueError) as e:
            msg = e.errors()[0]["msg"] if isinstance(e, ValidationError) else str(e)
            report.error(i, f"Некорректный прогноз: {msg}")
    known = known_channel_ids(db, {row["channel_id"] for _, row in clean})
    valid = []
    for i, row in clean:
        if row["channel_id"] in known:
            valid.append(row)
        else:
            report.error(i, f"Канал {row['channel_id']} не найден в справочнике")
    cols = ["prob", "health", "risk_level", "top_factors", "model_version"]
    report.rows_imported = upsert(db, Prediction.__table__, valid, ["channel_id", "at", "horizon_h"], cols)
    # Для SSE: критические прогнозы последнего загруженного момента
    if valid:
        last_at = max(r["at"] for r in valid)
        critical_ids = [r["channel_id"] for r in valid if r["at"] == last_at and r["risk_level"] == "critical"]
        if critical_ids:
            report.imported_predictions = list(
                db.scalars(select(Prediction).where(Prediction.at == last_at, Prediction.channel_id.in_(critical_ids)))
            )


IMPORTERS = {
    "objects": _import_objects,
    "channels": _import_channels,
    "events": _import_events,
    "predictions": _import_predictions,
}


def import_file(db: Session, kind: str, file_name: str, content: bytes) -> ImportReport:
    report = ImportReport(kind=kind, file_name=file_name)
    rows = map_columns(read_table(content, file_name), kind)
    report.rows_total = len(rows)
    IMPORTERS[kind](db, rows, report)
    return report
