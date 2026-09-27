import io
from datetime import datetime, timedelta
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import Select, func, or_, select
from sqlalchemy.orm import Session

from app import schemas, xmlio
from app.api.predictions import load_predictions
from app.access import Access, get_access, require
from app.audit import audit
from app.db import get_db
from app.deps import Page, at_param, page_param
from app.labels import ALERT_LEVELS, DECISION_LABELS, RISK_LABELS, WORK_ORDER_STATUS_LABELS
from app.models import Channel, Decision, Prediction, PredictionOutcome, User
from app.recommendations.engine import advise_predictions
from app.risk import level_in
from app.security import get_current_user
from app.services import ObjectIndex, apply_channel_filters, prediction_items, snapshot_at, work_orders_by_prediction
from app.timeutil import sim_moment

router = APIRouter(tags=["Журнал прогнозов"])

EXPORT_MAX_ROWS = 50_000
# Горизонт прогноза от ML — всегда 24 ч (ТЗ бэкенда)
HORIZON_H = 24


class JournalFilters:
    """Фильтры журнала — общие для списка и выгрузки в XLSX."""

    def __init__(
        self,
        date_from: datetime | None = Query(
            None, description="Начало периода; по умолчанию 7 дней до последнего среза прогнозов не позже at"
        ),
        date_to: datetime | None = Query(None, description="Конец периода; по умолчанию at"),
        risk_level: list[schemas.RiskLevel] | None = Query(None, description="Можно несколько"),
        decision: Literal["none", "dispatch", "false_alarm", "monitor"] | None = Query(
            None, description="none — решения ещё нет"
        ),
        outcome: Literal["happened", "not_happened", "unknown"] | None = Query(None, description="Исход прогноза"),
        object_id: int | None = Query(None, description="Объект; включает все вложенные"),
        system_type: str | None = None,
        sensor_type: str | None = None,
        q: str | None = Query(None, description="Поиск по названию, тегу или id датчика"),
        only_alerts: bool = Query(
            True, description="Только тревожные прогнозы (риск и критично) и прогнозы с решением диспетчера"
        ),
    ):
        self.date_from, self.date_to = date_from, date_to
        self.risk_level, self.decision, self.outcome = risk_level, decision, outcome
        self.object_id, self.system_type, self.sensor_type, self.q = object_id, system_type, sensor_type, q
        self.only_alerts = only_alerts


def journal_query(db: Session, idx: ObjectIndex, at: datetime, f: JournalFilters) -> Select:
    """Запрос id прогнозов журнала, от новых к старым."""
    date_to = min(f.date_to, at) if f.date_to else at
    # Период по умолчанию — от момента данных: «сейчас» после конца журнала заказчика не даёт пустой журнал
    date_from = f.date_from or (snapshot_at(db, at) or at) - timedelta(days=7)
    if date_from > date_to:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "Начало периода позже конца")
    # Актуальное решение по прогнозу на момент at — последнее по id
    latest = (
        select(Decision.prediction_id, func.max(Decision.id).label("decision_id"))
        .where(Decision.created_at <= at)
        .group_by(Decision.prediction_id)
        .subquery()
    )
    # Соединения с решениями и исходами — только когда по ним фильтруют: на реальных данных за 7 дней ≈ 1,3 млн
    # почасовых прогнозов, и лишний hash join с 10 млн исходов стоил больше самого журнала (нагрузочный тест)
    stmt = (
        select(Prediction.id)
        .join(Channel, Channel.id == Prediction.channel_id)
        .where(Prediction.at >= date_from, Prediction.at <= date_to)
    )
    if f.decision:
        stmt = stmt.outerjoin(latest, latest.c.prediction_id == Prediction.id)
        stmt = stmt.outerjoin(Decision, Decision.id == latest.c.decision_id)
    if f.outcome:
        stmt = stmt.outerjoin(PredictionOutcome, PredictionOutcome.prediction_id == Prediction.id)
    stmt = apply_channel_filters(stmt, idx, f.object_id, f.system_type, f.sensor_type, f.q)
    if f.only_alerts:
        # «Тревожные или с решением». Прогнозы с решением — отдельным маленьким запросом и списком id: так условие
        # остаётся индексным (BitmapOr по ix_predictions_at_risk и первичному ключу), а не OR с подзапросом,
        # из-за которого PostgreSQL читал все прогнозы подряд (0,6 с на каждый запрос журнала)
        decided = list(db.scalars(
            select(Decision.prediction_id).distinct()
            .join(Prediction, Prediction.id == Decision.prediction_id)
            .where(Decision.created_at <= at, Prediction.at >= date_from, Prediction.at <= date_to)
        ))
        stmt = stmt.where(or_(level_in(ALERT_LEVELS), Prediction.id.in_(decided)) if decided else level_in(ALERT_LEVELS))
    if f.risk_level:
        stmt = stmt.where(level_in(f.risk_level))
    if f.decision == "none":
        stmt = stmt.where(latest.c.decision_id.is_(None))
    elif f.decision:
        stmt = stmt.where(Decision.decision_type == f.decision)
    # Симуляция: исход известен, только если горизонт (24 ч) закрылся к модельному моменту
    moment = sim_moment.get()
    closed = Prediction.at <= moment - timedelta(hours=HORIZON_H) if moment is not None else None
    if f.outcome == "happened":
        stmt = stmt.where(PredictionOutcome.happened.is_(True))
        stmt = stmt.where(closed) if closed is not None else stmt
    elif f.outcome == "not_happened":
        stmt = stmt.where(PredictionOutcome.happened.is_(False))
        stmt = stmt.where(closed) if closed is not None else stmt
    elif f.outcome == "unknown":
        unknown = PredictionOutcome.prediction_id.is_(None)
        stmt = stmt.where(or_(unknown, ~closed) if closed is not None else unknown)
    return stmt.order_by(Prediction.at.desc(), Prediction.prob.desc(), Prediction.id)


def journal_items(db: Session, idx: ObjectIndex, ids: list[int], at: datetime) -> list[schemas.JournalItem]:
    preds = load_predictions(db, ids)
    base_items = prediction_items(db, preds, idx, at)
    orders = work_orders_by_prediction(db, ids, at)
    advice = advise_predictions(db, preds)
    return [
        schemas.JournalItem(
            **item.model_dump(),
            at=p.at,
            work_order_status=orders[p.id][1] if p.id in orders else None,
            work_order_number=orders[p.id][2] if p.id in orders else None,
            recommendation=schemas.JournalRecommendation(
                rule_id=advice[p.id].rule_id, title=advice[p.id].title, action=advice[p.id].action,
                priority=advice[p.id].priority, fault_kind_label=advice[p.id].fault_kind_label,
            ),
        )
        for item, p in zip(base_items, preds)
    ]


@router.get(
    "/api/journal",
    response_model=schemas.JournalList,
    summary="Журнал прогнозов с решениями, исходами и заявками",
    dependencies=[Depends(get_current_user)],
)
def journal(
    f: JournalFilters = Depends(),
    page: Page = Depends(page_param),
    at=Depends(at_param),
    access: Access = Depends(get_access),
    db: Session = Depends(get_db),
):
    idx = ObjectIndex.load(db, access)
    stmt = journal_query(db, idx, at, f)
    total = db.scalar(select(func.count()).select_from(stmt.order_by(None).subquery()))
    ids = db.scalars(stmt.limit(page.limit).offset(page.offset)).all()
    return schemas.JournalList(at=at, total=total, items=journal_items(db, idx, list(ids), at))


# Строки, которые Excel/LibreOffice исполнят как формулу (CSV/XLSX-инъекция, OWASP). Комментарии и причины
# вводят люди, названия датчиков приходят из импорта — такие значения выгружаем с апострофом, как текст
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n", "\uff1d", "\uff0b", "\uff0d", "\uff20")


def excel_safe(value):
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


EXPORT_COLUMNS = [
    ("Время прогноза", 20), ("ID прогноза", 12), ("ID датчика", 12), ("Датчик", 28), ("Тип датчика", 24),
    ("Система", 22), ("Объект", 40), ("Вероятность отказа", 12), ("Индекс здоровья", 10), ("Уровень риска", 14),
    ("Причины", 70), ("Причины (формулировки модели)", 60), ("Решение", 20), ("Причина решения", 30), ("Комментарий", 30), ("Кто решил", 18),
    ("Исход", 14), ("Время отказа", 20), ("Заявка", 18), ("Статус заявки", 14), ("Рекомендация", 70),
]


def journal_xml(db: Session, idx: ObjectIndex, ids: list[int], at: datetime) -> bytes:
    """Журнал в XML: <journal at rows><prediction id>…</prediction></journal> — поля как в колонках XLSX."""
    records = []
    for start in range(0, len(ids), 1000):
        for it in journal_items(db, idx, ids[start : start + 1000], at):
            d, o, rec = it.decision, it.outcome, it.recommendation
            records.append(("prediction", {"id": it.prediction_id}, {
                "at": it.at, "channel_id": it.channel.id, "channel": it.channel.name, "sensor_type": it.channel.sensor_type,
                "system_type": it.channel.system_type, "object_id": it.object.id, "object": " → ".join(it.object.path),
                "prob": round(it.prob, 4), "health": it.health, "risk_level": it.risk_level,
                "factors": "; ".join(f.text for f in it.factors_human), "factors_model": "; ".join(it.factors),
                "decision": d.decision_type if d else None, "decision_reason": d.reason.name if d else None,
                "decision_comment": d.comment if d else None, "decision_user": d.user.full_name if d and d.user else None,
                "outcome": None if o is None else ("happened" if o.happened else "not_happened"),
                "fault_at": o.fault_at if o else None, "work_order": it.work_order_number,
                "work_order_status": it.work_order_status,
                "recommendation_rule": rec.rule_id if rec else None,
                "recommendation": f"{rec.title}: {rec.action}" if rec else None,
            }))
    return xmlio.build("journal", {"at": at, "rows": len(records)}, records)


@router.get(
    "/api/export/journal",
    summary="Выгрузка журнала в XLSX или XML (те же фильтры, что у /api/journal)",
    response_class=StreamingResponse,
    responses={200: {"content": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {},
                                 "application/xml": {}},
                     "description": "Файл XLSX (по умолчанию) или XML (`format=xml`)"}},
)
def export_journal(
    request: Request,
    f: JournalFilters = Depends(),
    format: Literal["xlsx", "xml"] = Query("xlsx", description="xlsx — таблица Excel, xml — для обмена с другими системами"),
    at=Depends(at_param),
    access: Access = Depends(require("export")),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    idx = ObjectIndex.load(db, access)
    ids = list(db.scalars(journal_query(db, idx, at, f).limit(EXPORT_MAX_ROWS)))
    if format == "xml":
        content = journal_xml(db, idx, ids, at)
        audit(db, request, user, "export_journal", "journal",
              details={"rows": len(ids), "at": at.isoformat(), "format": "xml"})
        db.commit()
        return StreamingResponse(
            io.BytesIO(content), media_type="application/xml",
            headers={"Content-Disposition": f'attachment; filename="journal_{at:%Y%m%d_%H%M}.xml"'},
        )

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Журнал прогнозов")
    for i, (_, width) in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    header = []
    for title, _ in EXPORT_COLUMNS:
        cell = WriteOnlyCell(ws, value=title)
        cell.font = Font(bold=True)
        header.append(cell)
    ws.append(header)
    for start in range(0, len(ids), 1000):
        for it in journal_items(db, idx, ids[start : start + 1000], at):
            d, o = it.decision, it.outcome
            ws.append([excel_safe(v) for v in [
                it.at, it.prediction_id, it.channel.id, it.channel.name, it.channel.sensor_type,
                it.channel.system_type, " → ".join(it.object.path), round(it.prob, 3), it.health,
                RISK_LABELS[it.risk_level][0], "; ".join(f.text for f in it.factors_human), "; ".join(it.factors),
                DECISION_LABELS[d.decision_type] if d else "", d.reason.name if d else "", (d.comment or "") if d else "",
                d.user.full_name if d and d.user else "",
                "" if o is None else ("Сбылся" if o.happened else "Не сбылся"),
                o.fault_at if o and o.fault_at else None, it.work_order_number or it.work_order_id,
                WORK_ORDER_STATUS_LABELS[it.work_order_status] if it.work_order_status else "",
                f"{it.recommendation.title}: {it.recommendation.action}" if it.recommendation else "",
            ]])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    audit(db, request, user, "export_journal", "journal", details={"rows": len(ids), "at": at.isoformat()})
    db.commit()
    file_name = f"journal_{at:%Y%m%d_%H%M}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{file_name}"'},
    )
