"""Журнал тревожных событий (ответ заказчика: «в журнал попадают любые тревожные события»; общий ТЗ, приложение 2 —
журнал технологических событий). Источник — оперативный контур events_recent (последние 30 дней и входящий поток).
Контекст — прозрачные правила app/event_context.py: вероятная плановая проверка, групповое отключение, сбой значения."""

import io
from datetime import datetime, timedelta
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import Select, func, or_, select
from sqlalchemy.orm import Session

from app import event_context, schemas
from app.access import Access, get_access, require
from app.api.journal import excel_safe
from app.audit import audit
from app.db import get_db
from app.deps import Page, at_param, page_param
from app.geo.pickets import parse as parse_picket
from app.models import Channel, EventRecent, User
from app.security import get_current_user
from app.services import ObjectIndex, apply_channel_filters, channel_ref

router = APIRouter(tags=["Журнал событий"])

MAX_PERIOD = timedelta(days=31)
EXPORT_MAX_ROWS = 50_000
# Только события статусов (смены состояния) — без потока числовых показаний, кроме служебных кодов
STATUS_ONLY = EventRecent.value_text.is_not(None)


class EventFilters:
    def __init__(
        self,
        date_from: datetime | None = Query(None, description="Начало периода; по умолчанию сутки до конца периода"),
        date_to: datetime | None = Query(None, description="Конец периода; по умолчанию at (не позже последнего события)"),
        object_id: int | None = Query(None, description="Объект; включает вложенные"),
        channel_id: int | None = Query(None, description="Один датчик (блок событий на карточке)"),
        system_type: str | None = None,
        sensor_type: str | None = None,
        q: str | None = Query(None, description="Поиск по названию, тегу или id датчика"),
        only_alarms: bool = Query(True, description="Только события с флагом «тревожное»"),
        kind: Literal["all", "status", "values"] = Query(
            "status", description="status — сообщения о состоянии (смены статуса), values — только показания, all — всё"
        ),
        tag: Literal["planned_check", "group_off", "value_failure"] | None = Query(
            None, description="Только события с этим контекстом (проверяется на странице результата)"
        ),
    ):
        self.date_from, self.date_to = date_from, date_to
        self.object_id, self.channel_id = object_id, channel_id
        self.system_type, self.sensor_type, self.q = system_type, sensor_type, q
        self.only_alarms, self.kind, self.tag = only_alarms, kind, tag


def event_query(db: Session, idx: ObjectIndex, at: datetime, f: EventFilters) -> tuple[Select, datetime, datetime]:
    date_to = min(f.date_to, at) if f.date_to else at
    # «Сейчас» после конца данных: период — от последнего события не позже at, как у журнала прогнозов
    last = db.scalar(select(func.max(EventRecent.ts)).where(EventRecent.ts <= date_to))
    if f.date_to is None and last is not None:
        date_to = last
    date_from = f.date_from or date_to - timedelta(days=1)
    if date_from > date_to:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "Начало периода позже конца")
    if date_to - date_from > MAX_PERIOD:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_CONTENT, "Период не может быть длиннее 31 дня")
    stmt = (
        select(EventRecent.id)
        .join(Channel, Channel.id == EventRecent.channel_id)
        .where(EventRecent.ts > date_from, EventRecent.ts <= date_to)
    )
    stmt = apply_channel_filters(stmt, idx, f.object_id, f.system_type, f.sensor_type, f.q)
    if f.channel_id is not None:
        stmt = stmt.where(EventRecent.channel_id == f.channel_id)
    if f.only_alarms:
        stmt = stmt.where(EventRecent.is_alarm.is_(True))
    if f.kind == "status":
        # Статусы и нечисловые значения (в том числе «01.01.1970» — сбой значения); числовые — только служебные коды
        stmt = stmt.where(or_(STATUS_ONLY, EventRecent.value_num.is_(None),
                              EventRecent.value_num.in_(event_context.SENTINEL_CODES)))
    elif f.kind == "values":
        stmt = stmt.where(EventRecent.value_num.is_not(None))
    return stmt.order_by(EventRecent.ts.desc(), EventRecent.id.desc()), date_from, date_to


def event_items(db: Session, idx: ObjectIndex, ids: list[int]) -> list[schemas.EventItem]:
    if not ids:
        return []
    events = {e.id: e for e in db.scalars(select(EventRecent).where(EventRecent.id.in_(ids)))}
    ordered = [events[i] for i in ids if i in events]
    channels = {c.id: c for c in db.scalars(select(Channel).where(Channel.id.in_({e.channel_id for e in ordered})))}
    parent_of = {i: n.parent_id for i, n in idx.nodes.items()}
    tags = event_context.annotate(db, ordered, channels, parent_of)
    items = []
    for e in ordered:
        ch = channels[e.channel_id]
        kind = event_context.value_kind(e)
        items.append(schemas.EventItem(
            id=e.id, ts=e.ts, channel=channel_ref(ch), object=idx.ref(ch.object_id),
            picket=parse_picket(ch.name).label or None, value=e.value_raw, value_num=e.value_num,
            value_kind=kind, value_kind_label=event_context.VALUE_KIND_LABELS[kind], is_alarm=e.is_alarm,
            tags=[schemas.EventTag(code=t.code, label=t.label, reason=t.reason) for t in tags[e.id]],
        ))
    return items


@router.get(
    "/api/events",
    response_model=schemas.EventList,
    summary="Журнал тревожных событий и смен статуса с контекстом (плановая проверка, групповое отключение, сбой)",
    description=(
        "События оперативного контура (`events_recent`): по умолчанию — тревожные сообщения о состоянии за сутки до "
        "`at`. Контекст — прозрачные правила, не ML: «Вероятная плановая проверка» (газ в будни 9–14), «Групповое "
        "отключение — вероятно, плановые работы» (≥ 80 % и не меньше 5 однотипных датчиков объекта отключились в "
        "±12 ч), «Сбой значения» (служебный код, дата вместо показания, отрицательный газ). Только датчики области "
        "пользователя. `tag` фильтрует страницу результата (правила считаются на странице)."
    ),
    dependencies=[Depends(get_current_user)],
)
def events(
    f: EventFilters = Depends(),
    page: Page = Depends(page_param),
    at=Depends(at_param),
    access: Access = Depends(get_access),
    db: Session = Depends(get_db),
):
    idx = ObjectIndex.load(db, access)
    stmt, date_from, date_to = event_query(db, idx, at, f)
    total = db.scalar(select(func.count()).select_from(stmt.order_by(None).subquery()))
    ids = list(db.scalars(stmt.limit(page.limit).offset(page.offset)))
    items = event_items(db, idx, ids)
    if f.tag:
        items = [i for i in items if any(t.code == f.tag for t in i.tags)]
    return schemas.EventList(at=at, date_from=date_from, date_to=date_to, total=total, items=items)


EXPORT_COLUMNS = [("Время", 20), ("ID датчика", 12), ("Датчик", 30), ("Тип датчика", 22), ("Система", 22),
                  ("Объект", 40), ("Пикет", 18), ("Значение", 22), ("Вид значения", 18), ("Тревожное", 10),
                  ("Контекст", 40), ("Почему", 70)]


@router.get(
    "/api/export/events",
    summary="Выгрузка журнала событий в XLSX (те же фильтры, что у /api/events)",
    response_class=StreamingResponse,
    responses={200: {"content": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {}},
                     "description": "Файл XLSX, до 50 000 строк"}},
)
def export_events(
    request: Request,
    f: EventFilters = Depends(),
    at=Depends(at_param),
    access: Access = Depends(require("export")),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    idx = ObjectIndex.load(db, access)
    stmt, date_from, date_to = event_query(db, idx, at, f)
    ids = list(db.scalars(stmt.limit(EXPORT_MAX_ROWS)))
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Журнал событий")
    for i, (_, width) in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    header = []
    for title, _ in EXPORT_COLUMNS:
        cell = WriteOnlyCell(ws, value=title)
        cell.font = Font(bold=True)
        header.append(cell)
    ws.append(header)
    rows = 0
    for start in range(0, len(ids), 1000):
        for it in event_items(db, idx, ids[start : start + 1000]):
            if f.tag and not any(t.code == f.tag for t in it.tags):
                continue
            rows += 1
            ws.append([excel_safe(v) for v in [
                it.ts, it.channel.id, it.channel.name, it.channel.sensor_type, it.channel.system_type,
                " → ".join(it.object.path), it.picket or "", it.value, it.value_kind_label, "да" if it.is_alarm else "",
                "; ".join(t.label for t in it.tags), "; ".join(t.reason for t in it.tags),
            ]])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    audit(db, request, user, "export_events", "events",
          details={"rows": rows, "at": at.isoformat(), "from": date_from.isoformat(), "to": date_to.isoformat()})
    db.commit()
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="events_{date_to:%Y%m%d_%H%M}.xlsx"'},
    )
