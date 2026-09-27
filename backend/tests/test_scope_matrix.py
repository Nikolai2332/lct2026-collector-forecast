"""Матрица «роль × эндпоинт × чужой объект» по всей схеме OpenAPI (аудит 2, docs/SECURITY_AUDIT.md).

Каждая операция API должна быть здесь классифицирована: новый эндпоинт без строки в OPERATIONS роняет
test_every_operation_is_classified — так область видимости и права нельзя «забыть».

- Чтение: пользователь с областью (техник — один комплекс; временный диспетчер с той же областью) получает
  только свои датчики и объекты. Ответ проверяется целиком: все channel_id, channel.id, channel_ids, object_id,
  object.id в JSON, id датчиков в XLSX/XML/CSV. Чужой id в пути — 404; чужой object_id в фильтре — 404;
  чужой channel_id / prediction_id в фильтре — пустой результат.
- Запись: чужой прогноз, датчик или заявка — 404 (или «не найден» в пакетной операции) для роли, у которой
  право есть; роль без права — 403 на каждой пишущей операции.
"""

import csv
import io
import xml.etree.ElementTree as ET
from datetime import datetime

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.access import PERMISSIONS
from app.main import app
from app.models import Channel, Decision, Object, Prediction, User, WorkOrder
from app.services import ObjectIndex
from app.access import load_access

AT = "2026-08-01T12:00:00"

# Операция → как проверяется. read — ответ сканируется; id — чужой id в пути; public/self — без данных объектов;
# write:<право> — роль без права получает 403, чужие объекты — 404; admin:<право> — только администратор.
OPERATIONS: dict[tuple[str, str], str] = {
    ("GET", "/api/health"): "public",
    ("POST", "/api/auth/login"): "public",
    ("POST", "/api/auth/token"): "public",
    ("GET", "/api/auth/me"): "self",
    ("GET", "/api/reasons"): "dictionary",
    ("GET", "/api/recommendations"): "dictionary",
    ("GET", "/api/dictionaries"): "read",
    ("GET", "/api/dashboard/summary"): "read",
    ("GET", "/api/objects"): "read",
    ("GET", "/api/objects/{object_id}"): "id",
    ("GET", "/api/channels"): "read",
    ("GET", "/api/channels/{channel_id}"): "id",
    ("GET", "/api/channels/{channel_id}/history"): "id",
    ("GET", "/api/predictions"): "read",
    ("GET", "/api/predictions/{prediction_id}"): "id",
    ("GET", "/api/predictions/{prediction_id}/recommendation"): "id",
    ("POST", "/api/predictions/{prediction_id}/decision"): "write:decide",
    ("GET", "/api/journal"): "read",
    ("GET", "/api/export/journal"): "export",
    ("POST", "/api/work-orders"): "write:work_orders",
    ("GET", "/api/work-orders"): "read",
    ("GET", "/api/work-orders/{work_order_id}"): "id",
    ("PATCH", "/api/work-orders/{work_order_id}"): "write:work_order_status",
    ("DELETE", "/api/work-orders/{work_order_id}"): "write:work_orders",
    ("GET", "/api/maintenance/plan"): "read",
    ("POST", "/api/maintenance/drafts"): "write:work_orders",
    ("GET", "/api/model/metrics"): "read",
    ("GET", "/api/model/thresholds"): "model",
    ("POST", "/api/ingest/events"): "write:data_import",
    ("POST", "/api/import"): "write:data_import",
    ("GET", "/api/notifications"): "read",
    ("POST", "/api/notifications/ticket"): "self",
    ("GET", "/api/notifications/stream"): "sse",
    ("POST", "/api/notifications/test"): "write:data_import",
    ("GET", "/api/sim/state"): "self",
    ("POST", "/api/sim/start"): "write:sim_control",
    ("POST", "/api/sim/stop"): "write:sim_control",
    ("GET", "/api/settings"): "self",
    ("PUT", "/api/settings"): "write:settings",
    ("GET", "/api/users"): "write:users_admin",
    ("PUT", "/api/users/{user_id}/access"): "write:users_admin",
    ("GET", "/api/geo/collectors"): "read",
    ("GET", "/api/geo/channels"): "read",
    ("GET", "/api/events"): "read",
    ("GET", "/api/export/events"): "export",
    ("GET", "/api/export/decisions"): "export",
}


def test_every_operation_is_classified():
    ops = {(m.upper(), p) for p, item in app.openapi()["paths"].items() for m in item}
    assert ops - set(OPERATIONS) == set(), "новый эндпоинт: добавьте его в матрицу областей видимости"
    assert set(OPERATIONS) - ops == set(), "эндпоинт удалён — уберите его из матрицы"


# ---------- Данные: область техника и чужие объекты ----------


@pytest.fixture(scope="module")
def world(session_factory, client):
    token = client.post("/api/auth/login", json={"username": "technician", "password": "technician123"}).json()
    with session_factory() as s:
        user = s.scalar(select(User).where(User.username == "technician"))
        idx = ObjectIndex.load(s, load_access(s, user))
        inside = set(s.scalars(select(Channel.id).where(Channel.object_id.in_(idx.channel_objects))))
        outside = set(s.scalars(select(Channel.id))) - inside
        visible_objects = set(idx.nodes)
        out_objects = set(s.scalars(select(Object.id))) - visible_objects
        at = datetime.fromisoformat(AT)
        out_pred = s.scalar(select(Prediction.id).where(Prediction.channel_id.in_(outside), Prediction.at <= at)
                            .order_by(Prediction.at.desc()).limit(1))
        in_pred = s.scalar(select(Prediction.id).where(Prediction.channel_id.in_(inside), Prediction.at <= at)
                           .order_by(Prediction.at.desc()).limit(1))
        out_wo = s.scalar(select(WorkOrder.id).where(WorkOrder.channel_id.in_(outside)).limit(1))
        out_name = s.scalar(select(Channel.name).where(Channel.id == min(outside)))
        has_out_decisions = s.scalar(select(Decision.id).join(Prediction, Prediction.id == Decision.prediction_id)
                                     .where(Prediction.channel_id.in_(outside)).limit(1)) is not None
    assert inside and outside and out_objects and out_pred and in_pred and out_wo
    return {
        "tech": {"Authorization": f"Bearer {token['access_token']}"},
        "complex_id": token["user"]["scope"][0]["id"],
        "inside": inside, "outside": outside, "visible_objects": visible_objects, "out_objects": out_objects,
        "out_object": min(out_objects), "out_channel": min(outside), "out_pred": out_pred, "in_pred": in_pred,
        "out_wo": out_wo, "out_name": out_name, "has_out_decisions": has_out_decisions,
    }


@pytest.fixture(scope="module")
def scoped_writer(client, session_factory, world):
    """Диспетчер района с областью «один комплекс»: права на решения и заявки есть, чужие объекты — нет."""
    from app.bootstrap import upsert_user

    with session_factory() as s:
        u = upsert_user(s, "tmp_matrix", "tmp-matrix-pass-123", "dispatcher", "Матрица", [world["complex_id"]])
        s.commit()
        uid = u.id
    token = client.post("/api/auth/login", json={"username": "tmp_matrix", "password": "tmp-matrix-pass-123"}).json()
    yield {"Authorization": f"Bearer {token['access_token']}"}
    with session_factory() as s:
        s.query(WorkOrder).filter(WorkOrder.created_by == uid).delete()
        s.query(Decision).filter(Decision.user_id == uid).delete()
        s.delete(s.get(User, uid))
        s.commit()


def _scan(data, found: dict[str, set[int]]) -> dict[str, set[int]]:
    """Все id датчиков и объектов в ответе, на любой глубине."""
    if isinstance(data, dict):
        if isinstance(data.get("id"), int) and {"sensor_type", "tag"} <= data.keys():  # датчик (ChannelListItem, ChannelRef)
            found["channels"].add(data["id"])
        if isinstance(data.get("id"), int) and {"level", "kind"} <= data.keys():  # узел дерева (ObjectNode)
            found["objects"].add(data["id"])
        for k, v in data.items():
            if k in ("channel_id",) and isinstance(v, int):
                found["channels"].add(v)
            elif k == "channel_ids" and isinstance(v, list):
                found["channels"].update(x for x in v if isinstance(x, int))
            elif k == "channel" and isinstance(v, dict) and isinstance(v.get("id"), int):
                found["channels"].add(v["id"])
            elif k in ("object_id", "parent_id") and isinstance(v, int):
                found["objects"].add(v)
            elif k == "object" and isinstance(v, dict) and isinstance(v.get("id"), int):
                found["objects"].add(v["id"])
            _scan(v, found)
    elif isinstance(data, list):
        for v in data:
            _scan(v, found)
    return found


def _assert_own(body, world, where: str):
    found = _scan(body, {"channels": set(), "objects": set()})
    assert not found["channels"] & world["outside"], (where, "чужие датчики", found["channels"] & world["outside"])
    assert not found["objects"] & world["out_objects"], (where, "чужие объекты", found["objects"] & world["out_objects"])
    return found


READ_PARAMS: dict[str, list[dict]] = {
    "/api/dictionaries": [{}],
    "/api/dashboard/summary": [{}],
    "/api/objects": [{}],
    "/api/channels": [{"limit": 500}, {"q": "{out_name}", "limit": 500}, {"sort": "name", "limit": 500}],
    "/api/predictions": [{"limit": 500}, {"min_prob": 0, "limit": 500}, {"q": "{out_name}"}],
    "/api/journal": [{"limit": 500, "only_alerts": False, "date_from": "2026-07-27T00:00:00"},
                     {"q": "{out_name}", "only_alerts": False}],
    "/api/work-orders": [{"limit": 500}, {"channel_id": "{out_channel}"}, {"prediction_id": "{out_pred}"},
                         {"q": "ЗН"}],
    "/api/maintenance/plan": [{"limit": 500}],
    "/api/model/metrics": [{"date_from": "2026-07-27", "date_to": "2026-08-01"}],
    "/api/notifications": [{"hours": 168, "limit": 500}],
    "/api/geo/collectors": [{}, {"format": "wkt"}],
    "/api/geo/channels": [{}, {"format": "wkt"}],
    "/api/events": [{"limit": 500, "only_alarms": False, "kind": "all", "date_from": "2026-07-27T00:00:00"},
                    {"channel_id": "{out_channel}", "only_alarms": False, "kind": "all"},
                    {"q": "{out_name}", "only_alarms": False, "kind": "all"}],
}
# Фильтр object_id с чужим объектом — 404 («Объект не найден»), а не пустой или чужой результат
OBJECT_FILTER = ["/api/channels", "/api/predictions", "/api/journal", "/api/work-orders", "/api/maintenance/plan",
                 "/api/geo/collectors", "/api/geo/channels", "/api/events", "/api/export/journal",
                 "/api/export/events"]


def _fill(params: dict, world) -> dict:
    return {k: (v.format(**world) if isinstance(v, str) and "{" in v else v) for k, v in params.items()}


def test_read_matrix_covers_all_read_operations():
    reads = {p for (m, p), kind in OPERATIONS.items() if kind == "read"}
    assert reads == set(READ_PARAMS)


@pytest.mark.parametrize("path", sorted(READ_PARAMS))
def test_scoped_reads_contain_only_own_objects(client, world, path):
    for params in READ_PARAMS[path]:
        r = client.get(path, params={"at": AT, **_fill(params, world)}, headers=world["tech"])
        assert r.status_code == 200, (path, params, r.status_code, r.text[:200])
        _assert_own(r.json(), world, f"{path} {params}")


def test_scoped_reads_are_not_empty(client, world):
    """Сканер не должен «проходить» на пустых ответах: у техника свои данные есть."""
    for path in ("/api/channels", "/api/predictions", "/api/geo/channels"):
        found = _scan(client.get(path, params={"at": AT, "limit": 500}, headers=world["tech"]).json(),
                      {"channels": set(), "objects": set()})
        assert found["channels"] and found["channels"] <= world["inside"], path
    # и видит чужое там, где оно есть: у ОДС тот же сканер находит датчики вне области техника
    ods = client.post("/api/auth/login", json={"username": "ods", "password": "ods123"}).json()["access_token"]
    body = client.get("/api/channels", params={"at": AT, "limit": 500}, headers={"Authorization": f"Bearer {ods}"}).json()
    assert _scan(body, {"channels": set(), "objects": set()})["channels"] & world["outside"]


@pytest.mark.parametrize("path", OBJECT_FILTER)
def test_foreign_object_filter_is_404(client, world, path):
    r = client.get(path, params={"at": AT, "object_id": world["out_object"]}, headers=world["tech"])
    assert r.status_code == 404, (path, r.status_code)


@pytest.mark.parametrize("path, key", [
    ("/api/objects/{object_id}", "out_object"),
    ("/api/channels/{channel_id}", "out_channel"),
    ("/api/channels/{channel_id}/history", "out_channel"),
    ("/api/predictions/{prediction_id}", "out_pred"),
    ("/api/predictions/{prediction_id}/recommendation", "out_pred"),
    ("/api/work-orders/{work_order_id}", "out_wo"),
])
def test_foreign_ids_are_404(client, world, path, key):
    url = path.replace(path[path.index("{"):path.index("}") + 1], str(world[key]))
    assert client.get(url, params={"at": AT}, headers=world["tech"]).status_code == 404
    # проверка «сквозь»: у ОДС тот же id есть
    ods = client.post("/api/auth/login", json={"username": "ods", "password": "ods123"}).json()["access_token"]
    assert client.get(url, params={"at": AT}, headers={"Authorization": f"Bearer {ods}"}).status_code == 200


def test_id_matrix_covers_all_id_operations():
    ids = {p for (m, p), kind in OPERATIONS.items() if kind == "id"}
    params = {p for p, _ in test_foreign_ids_are_404.pytestmark[0].args[1]}
    assert ids == params


def test_own_object_card_lists_only_own_channels(client, world):
    r = client.get(f"/api/objects/{world['complex_id']}", params={"at": AT, "limit": 500}, headers=world["tech"])
    assert r.status_code == 200
    _assert_own(r.json(), world, "object card")
    r = client.get(f"/api/predictions/{world['in_pred']}/recommendation", params={"at": AT}, headers=world["tech"])
    assert r.status_code == 200
    _assert_own(r.json(), world, "recommendation")


# ---------- Выгрузки ----------


def test_exports_contain_only_own_channels(client, world):
    tech = world["tech"]
    r = client.get("/api/export/journal", params={"at": AT, "only_alerts": False, "date_from": "2026-07-27T00:00:00"},
                   headers=tech)
    ids = {row[2] for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(min_row=2, values_only=True)}
    assert ids and ids <= world["inside"]
    r = client.get("/api/export/journal", params={"at": AT, "only_alerts": False, "format": "xml",
                                                  "date_from": "2026-07-27T00:00:00"}, headers=tech)
    xml_ids = {int(p.findtext("channel_id")) for p in ET.fromstring(r.content).iter("prediction")}
    assert xml_ids and xml_ids <= world["inside"]
    r = client.get("/api/export/events", params={"at": AT, "only_alarms": False, "kind": "all",
                                                 "date_from": "2026-07-27T00:00:00"}, headers=tech)
    assert r.status_code == 200
    ids = {row[1] for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(min_row=2, values_only=True)}
    assert ids <= world["inside"]
    r = client.get("/api/export/events", params={"at": AT, "channel_id": world["out_channel"], "only_alarms": False,
                                                 "kind": "all"}, headers=tech)
    assert r.status_code == 200 and load_workbook(io.BytesIO(r.content)).active.max_row == 1
    r = client.get("/api/export/decisions", params={"at": AT, "date_from": "2026-07-01T00:00:00",
                                                    "date_to": "2026-08-01T12:00:00"}, headers=tech)
    assert r.status_code == 200
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][2] == "channel_id"
    assert {int(x[2]) for x in rows[1:]} <= world["inside"]
    if world["has_out_decisions"]:
        ods = client.post("/api/auth/login", json={"username": "ods", "password": "ods123"}).json()["access_token"]
        r = client.get("/api/export/decisions", params={"at": AT, "date_from": "2026-07-01T00:00:00",
                                                        "date_to": "2026-08-01T12:00:00"},
                       headers={"Authorization": f"Bearer {ods}"})
        all_ids = {int(x[2]) for x in list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))[1:]}
        assert all_ids & world["outside"]


def test_model_metrics_cache_does_not_cross_scopes(client, world):
    ods = client.post("/api/auth/login", json={"username": "ods", "password": "ods123"}).json()["access_token"]
    params = {"at": AT, "date_from": "2026-07-27", "date_to": "2026-08-01"}
    full = client.get("/api/model/metrics", params=params, headers={"Authorization": f"Bearer {ods}"}).json()["daily"]
    mine = client.get("/api/model/metrics", params=params, headers=world["tech"]).json()["daily"]
    assert sum(d["predicted"] for d in mine) < sum(d["predicted"] for d in full)
    again = client.get("/api/model/metrics", params=params, headers={"Authorization": f"Bearer {ods}"}).json()["daily"]
    assert again == full


# ---------- Запись: чужие объекты и права ----------


def test_scoped_writer_cannot_touch_foreign_objects(client, world, scoped_writer):
    w = scoped_writer
    r = client.post(f"/api/predictions/{world['out_pred']}/decision", headers=w,
                    json={"decision_type": "monitor", "reason_id": 11})
    assert r.status_code == 404
    assert client.post("/api/work-orders", headers=w, json={"channel_id": world["out_channel"]}).status_code == 404
    assert client.post("/api/work-orders", headers=w, json={"prediction_id": world["out_pred"]}).status_code == 404
    assert client.patch(f"/api/work-orders/{world['out_wo']}", headers=w, json={"assignee": "x"}).status_code == 404
    assert client.delete(f"/api/work-orders/{world['out_wo']}", headers=w).status_code == 404
    r = client.post("/api/maintenance/drafts", headers=w, params={"at": AT}, json={"prediction_ids": [world["out_pred"]]})
    assert r.status_code == 201 and not r.json()["created"]
    assert r.json()["skipped"][0]["reason"] == "Прогноз не найден"
    # свои — можно
    r = client.post("/api/work-orders", headers=w, params={"at": AT}, json={"prediction_id": world["in_pred"]})
    assert r.status_code == 201, r.text
    assert client.delete(f"/api/work-orders/{r.json()['id']}", headers=w).status_code == 204


WRITE_CALLS = {
    ("POST", "/api/predictions/{prediction_id}/decision"): {"json": {"decision_type": "monitor", "reason_id": 11}},
    ("POST", "/api/work-orders"): {"json": {"channel_id": 1}},
    ("PATCH", "/api/work-orders/{work_order_id}"): {"json": {"status": "cancelled"}},
    ("DELETE", "/api/work-orders/{work_order_id}"): {},
    ("POST", "/api/maintenance/drafts"): {"json": {"prediction_ids": [1]}},
    ("POST", "/api/ingest/events"): {"json": {"events": []}},
    ("POST", "/api/import"): {"data": {"kind": "events"}, "files": {"file": ("e.csv", b"a\n1\n", "text/csv")}},
    ("POST", "/api/notifications/test"): {},
    ("POST", "/api/sim/start"): {"json": {"day": "2026-08-01", "speed": 60}},
    ("POST", "/api/sim/stop"): {},
    ("PUT", "/api/settings"): {"json": {"risk_attention": 0.2, "risk_risk": 0.5, "risk_critical": 0.8,
                                        "notify_cooldown_hours": 24, "notify_per_slice": 3, "notify_sim_per_slice": 5}},
    ("GET", "/api/users"): {},
    ("PUT", "/api/users/{user_id}/access"): {"json": {"roles": ["admin"], "scope_object_ids": []}},
}


def test_write_matrix_covers_all_write_operations():
    assert {op for op, kind in OPERATIONS.items() if kind.startswith("write:")} == set(WRITE_CALLS)


@pytest.fixture(scope="module")
def role_headers(client, auth, world):
    ods = client.post("/api/auth/login", json={"username": "ods", "password": "ods123"}).json()["access_token"]
    return {"dispatcher": auth("dispatcher"), "engineer": auth("engineer"), "manager": auth("manager"),
            "admin": auth("admin"), "technician": world["tech"], "dispatcher_ods": {"Authorization": f"Bearer {ods}"}}


# Роли, у которых права нет (у кого есть — проверяется функциональными тестами)
DENIED = [(role, op) for op in sorted(WRITE_CALLS) for role in sorted(PERMISSIONS)
          if OPERATIONS[op].split(":", 1)[1] not in PERMISSIONS[role]]


@pytest.mark.parametrize("role, op", DENIED)
def test_write_without_permission_is_403(client, role_headers, world, session_factory, role, op):
    method, path = op
    with session_factory() as s:
        me = s.scalar(select(User.id).where(User.username == ("ods" if role == "dispatcher_ods" else role)))
    url = (path.replace("{prediction_id}", str(world["in_pred"])).replace("{work_order_id}", str(world["out_wo"]))
           .replace("{user_id}", str(me)))
    r = client.request(method, url, headers=role_headers[role], params={"at": AT}, **WRITE_CALLS[op])
    assert r.status_code == 403, (role, method, path, r.status_code, r.text[:200])
