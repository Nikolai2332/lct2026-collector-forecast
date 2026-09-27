"""Ролевая модель заказчика и области видимости (docs/SECURITY.md): матрица «роль × действие», фильтрация
по области во всех ответах с данными датчиков, 404 на чужие id, кэш, выгрузки, SSE, сложение ролей,
управление ролями администратором.

Демо-сид (small): диспетчер района и руководитель — весь район; техник — один комплекс (с датчиками);
ОДС, инженер, администратор — всё."""

import asyncio
import io
from datetime import datetime
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import schemas
from app.access import FULL_ACCESS, PERMISSIONS, load_access
from app.models import AuditLog, Channel, Object, Prediction, User, WorkOrder
from app.notifications import Broker, critical_transitions, slice_notifications
from app.services import ObjectIndex

AT = "2026-08-01T12:00:00"


@pytest.fixture(scope="module")
def tech_token(client):
    r = client.post("/api/auth/login", json={"username": "technician", "password": "technician123"})
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture(scope="module")
def ods_token(client):
    r = client.post("/api/auth/login", json={"username": "ods", "password": "ods123"})
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture(scope="module")
def h(auth, tech_token, ods_token):
    def headers(role: str) -> dict[str, str]:
        if role == "technician":
            return {"Authorization": f"Bearer {tech_token['access_token']}"}
        if role == "dispatcher_ods":
            return {"Authorization": f"Bearer {ods_token['access_token']}"}
        return auth(role)

    return headers


@pytest.fixture(scope="module")
def scope(session_factory, tech_token):
    """Комплекс техника, его объекты и датчики; датчики вне области."""
    complex_id = tech_token["user"]["scope"][0]["id"]
    with session_factory() as s:
        user = s.scalar(select(User).where(User.username == "technician"))
        idx = ObjectIndex.load(s, load_access(s, user))
        inside = set(s.scalars(select(Channel.id).where(Channel.object_id.in_(idx.channel_objects))))
        outside = set(s.scalars(select(Channel.id))) - inside
        objects_outside = set(s.scalars(select(Object.id))) - set(idx.nodes)
    assert inside and outside
    return {"complex_id": complex_id, "inside": inside, "outside": outside, "objects_outside": objects_outside,
            "objects": set(idx.channel_objects)}


def _pred_outside(session_factory, scope) -> int:
    with session_factory() as s:
        return s.scalar(select(Prediction.id).where(Prediction.channel_id.in_(scope["outside"]),
                                                    Prediction.at <= datetime.fromisoformat(AT))
                        .order_by(Prediction.at.desc()).limit(1))


# ---------- Вход, шапка, права ----------


def test_demo_users_have_customer_roles_and_scopes(tech_token, ods_token, client, auth):
    tech = tech_token["user"]
    assert tech["roles"] == ["technician"] and tech["unrestricted"] is False
    assert len(tech["scope"]) == 1 and tech["scope"][0]["level"] == 2 and tech["scope_label"] == tech["scope"][0]["name"]
    assert set(tech["permissions"]) == PERMISSIONS["technician"]
    ods = ods_token["user"]
    assert ods["roles"] == ["dispatcher_ods"] and ods["unrestricted"] is True and ods["scope_label"] == "все объекты"
    disp = client.get("/api/auth/me", headers=auth("dispatcher")).json()
    assert disp["roles"] == ["dispatcher"] and disp["scope"][0]["level"] == 1 and disp["role_label"] == "Диспетчер района"


# Матрица «роль × действие»: ожидаемый код — 403 (нет права) или «не 403» (право есть; дальше — 2xx/404/409/422)
ACTIONS = {
    "decide": ("POST", "/api/predictions/{pred}/decision", {"json": {"decision_type": "monitor", "reason_id": 11}}),
    "work_orders": ("POST", "/api/work-orders", {"json": {"channel_id": "{channel}", "priority": "low"}}),
    "data_import": ("POST", "/api/ingest/events", {"json": {"events": []}}),
    "sim_control": ("POST", "/api/sim/stop", {}),
    "settings": ("PUT", "/api/settings", {"json": {}}),
    "users_admin": ("GET", "/api/users", {}),
}
ROLES = ("dispatcher_ods", "dispatcher", "technician", "engineer", "manager", "admin")


@pytest.mark.parametrize("role", ROLES)
@pytest.mark.parametrize("permission", list(ACTIONS))
def test_permission_matrix(client, h, role, permission, session_factory, scope):
    method, path, kw = ACTIONS[permission]
    with session_factory() as s:
        pred = s.scalar(select(Prediction.id).where(Prediction.channel_id.in_(scope["inside"])).limit(1))
    channel = min(scope["inside"])
    path = path.format(pred=pred)
    kw = {k: ({kk: (channel if vv == "{channel}" else vv) for kk, vv in v.items()} if isinstance(v, dict) else v)
          for k, v in kw.items()}
    r = client.request(method, path, headers=h(role), **kw)
    allowed = permission in PERMISSIONS[role]
    assert (r.status_code != 403) == allowed, (role, permission, r.status_code, r.text[:200])
    if r.status_code == 201 and permission == "work_orders":
        with session_factory() as s:  # не оставляем черновик другим тестам
            s.delete(s.get(WorkOrder, r.json()["id"]))
            s.commit()


# ---------- Фильтрация по области ----------


def test_lists_are_filtered_by_scope(client, h, scope):
    tech, ods = h("technician"), h("dispatcher_ods")
    ch = client.get("/api/channels", params={"at": AT, "limit": 500}, headers=tech).json()
    assert ch["total"] == len(scope["inside"]) and {c["id"] for c in ch["items"]} == scope["inside"]
    assert client.get("/api/channels", params={"at": AT, "limit": 500}, headers=ods).json()["total"] > ch["total"]

    preds = client.get("/api/predictions", params={"at": AT, "limit": 500}, headers=tech).json()
    assert preds["items"] and {p["channel"]["id"] for p in preds["items"]} <= scope["inside"]

    journal = client.get("/api/journal", params={"at": AT, "limit": 500, "only_alerts": False,
                                                  "date_from": "2026-07-27T00:00:00"}, headers=tech).json()
    assert journal["items"] and {p["channel"]["id"] for p in journal["items"]} <= scope["inside"]

    tree = client.get("/api/objects", params={"at": AT}, headers=tech).json()["items"]
    assert len(tree) == 1 and [c["id"] for c in tree[0]["children"]] == [scope["complex_id"]]
    assert tree[0]["channels_count"] == len(scope["inside"])

    summary = client.get("/api/dashboard/summary", params={"at": AT}, headers=tech).json()
    assert summary["channels_total"] == len(scope["inside"])
    assert sum(summary["risk_distribution"].values()) <= len(scope["inside"])

    wo = client.get("/api/work-orders", params={"at": AT, "limit": 500}, headers=tech).json()
    assert {w["channel"]["id"] for w in wo["items"]} <= scope["inside"]
    all_wo = client.get("/api/work-orders", params={"at": AT, "limit": 500}, headers=ods).json()
    assert all_wo["total"] >= wo["total"]

    plan = client.get("/api/maintenance/plan", params={"at": AT}, headers=tech).json()
    assert all(i["channel"]["id"] in scope["inside"] for g in plan["groups"] for i in g["items"])

    feed = client.get("/api/notifications", params={"at": AT, "hours": 168}, headers=tech).json()
    assert all(n["prediction"]["channel"]["id"] in scope["inside"] for n in feed["items"])


def test_direct_ids_outside_scope_are_404(client, h, scope, session_factory):
    tech = h("technician")
    out_channel = min(scope["outside"])
    out_object = min(scope["objects_outside"])
    out_pred = _pred_outside(session_factory, scope)
    with session_factory() as s:
        out_wo = s.scalar(select(WorkOrder.id).where(WorkOrder.channel_id.in_(scope["outside"])).limit(1))
    assert out_wo is not None
    checks = [
        ("GET", f"/api/channels/{out_channel}", {}),
        ("GET", f"/api/channels/{out_channel}/history", {}),
        ("GET", f"/api/objects/{out_object}", {}),
        ("GET", f"/api/predictions/{out_pred}", {}),
        ("GET", f"/api/predictions/{out_pred}/recommendation", {}),
        ("GET", f"/api/work-orders/{out_wo}", {}),
        ("PATCH", f"/api/work-orders/{out_wo}", {"json": {"status": "cancelled"}}),
        ("GET", "/api/channels", {"params": {"object_id": out_object}}),
        ("GET", "/api/journal", {"params": {"object_id": out_object}}),
        ("GET", "/api/maintenance/plan", {"params": {"object_id": out_object}}),
    ]
    for method, path, kw in checks:
        r = client.request(method, path, headers=tech, **kw)
        assert r.status_code == 404, (method, path, r.status_code)
        assert r.json()["detail"] in ("Датчик не найден", "Объект не найден", "Прогноз не найден", "Заявка не найдена")
    # тот же id у ОДС — есть
    assert client.get(f"/api/channels/{out_channel}", headers=h("dispatcher_ods")).status_code == 200
    # у техника права на решение нет — 403 раньше проверки области
    r = client.post(f"/api/predictions/{out_pred}/decision", headers=tech, json={"decision_type": "monitor", "reason_id": 11})
    assert r.status_code == 403


def test_technician_changes_only_status(client, h, scope, session_factory, auth):
    r = client.post("/api/work-orders", headers=auth("dispatcher"), json={"channel_id": min(scope["inside"])})
    assert r.status_code == 201, r.text
    wo_id = r.json()["id"]
    try:
        tech = h("technician")
        assert client.patch(f"/api/work-orders/{wo_id}", headers=tech, json={"assignee": "Я"}).status_code == 403
        r = client.patch(f"/api/work-orders/{wo_id}", headers=tech, json={"status": "submitted"})
        assert r.status_code == 200 and r.json()["status"] == "submitted"
    finally:
        with session_factory() as s:
            s.delete(s.get(WorkOrder, wo_id))
            s.commit()


def test_exports_are_filtered_by_scope(client, h, scope):
    params = {"at": AT, "only_alerts": False, "date_from": "2026-07-31T00:00:00"}
    r = client.get("/api/export/journal", params=params, headers=h("technician"))
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    ids = {row[2] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert ids and ids <= scope["inside"]
    r = client.get("/api/export/journal", params={**params, "format": "xml"}, headers=h("technician"))
    root = ET.fromstring(r.content)
    xml_ids = {int(p.findtext("channel_id")) for p in root.iter("prediction")}
    assert xml_ids and xml_ids <= scope["inside"]


def test_dashboard_cache_is_per_scope(client, h):
    first = client.get("/api/dashboard/summary", params={"at": AT}, headers=h("dispatcher_ods")).json()
    tech = client.get("/api/dashboard/summary", params={"at": AT}, headers=h("technician")).json()
    again = client.get("/api/dashboard/summary", params={"at": AT}, headers=h("dispatcher_ods")).json()
    assert first == again
    assert tech["channels_total"] < first["channels_total"]
    a, b = tech["accuracy_30d"], first["accuracy_30d"]
    assert (a["true_positive"], a["false_positive"], a["false_negative"]) != (
        b["true_positive"], b["false_positive"], b["false_negative"]) or a["true_positive"] == 0
    quality_t = client.get("/api/model/metrics", params={"at": AT}, headers=h("technician")).json()["daily"]
    quality_o = client.get("/api/model/metrics", params={"at": AT}, headers=h("dispatcher_ods")).json()["daily"]
    assert sum(d["actual"] for d in quality_t) <= sum(d["actual"] for d in quality_o)


# ---------- SSE ----------


def test_sse_delivers_only_own_scope_including_summaries(session_factory, scope):
    """Брокер: у каждой области свои уведомления; сводки «ещё N» считаются по своей области."""
    with session_factory() as s:
        snaps = sorted(set(s.scalars(select(Prediction.at).where(Prediction.risk_level == "critical"))))

        async def run():
            b = Broker()
            b.bind_loop(asyncio.get_running_loop())
            q_tech = b.subscribe(frozenset(scope["objects"]))
            q_all = b.subscribe(None)
            for snap in snaps:
                b.publish_scoped(lambda sc, sn=snap: slice_notifications(
                    s, critical_transitions(s, sn, sn, sc), per_slice=1))
            got = {}
            for name, q in (("tech", q_tech), ("all", q_all)):
                items = []
                while not q.empty():
                    items.append(schemas.NotificationItem.model_validate_json(q.get_nowait()[1]))
                got[name] = items
            b.unsubscribe(q_tech)
            return got

        got = asyncio.run(run())
    assert got["all"], "в демо-сиде должны быть переходы в «Критично»"
    assert all(n.prediction.channel.id in scope["inside"] for n in got["tech"])
    tech_total = sum(n.count or 1 for n in got["tech"])
    with session_factory() as s:
        expected = sum(len(critical_transitions(s, sn, sn, frozenset(scope["objects"]))) for sn in snaps)
    assert tech_total == expected
    assert sum(n.count or 1 for n in got["all"]) >= tech_total


def test_single_publish_skips_foreign_subscribers(session_factory, scope):
    with session_factory() as s:
        pred = s.scalar(select(Prediction).where(Prediction.channel_id.in_(scope["outside"])).limit(1))
        from app.notifications import notifications_for

        item = notifications_for(s, [pred])[0]

    async def run():
        b = Broker()
        b.bind_loop(asyncio.get_running_loop())
        q_tech, q_all = b.subscribe(frozenset(scope["objects"])), b.subscribe(None)
        b.publish(item)
        return q_tech.qsize(), q_all.qsize()

    assert asyncio.run(run()) == (0, 1)


def test_stream_initial_snapshot_is_scoped(client, h, scope, monkeypatch):
    """Первый пакет потока (snapshot) — по области подписчика."""
    from app.api import notifications as api_n

    captured = {}
    real = api_n.slice_notifications

    def spy(db, transitions, moment=None, per_slice=None):
        items = real(db, transitions, moment, per_slice)
        captured["items"] = items
        raise RuntimeError("stop")  # поток бесконечный: останавливаемся на первом пакете

    monkeypatch.setattr(api_n, "slice_notifications", spy)
    with pytest.raises(RuntimeError):
        client.get("/api/notifications/stream", params={"at": AT}, headers=h("technician"))
    assert all(n.prediction.channel.id in scope["inside"] for n in captured["items"])


# ---------- Сложение ролей и управление администратором ----------


@pytest.fixture
def temp_user(session_factory):
    from app.bootstrap import upsert_user

    with session_factory() as s:
        u = upsert_user(s, "tmp_scoped", "tmp-password-123", "technician", "Временный", [])
        s.commit()
        uid = u.id
    yield uid
    with session_factory() as s:
        s.delete(s.get(User, uid))
        s.commit()


def test_scoped_role_without_scope_sees_nothing(client, temp_user):
    token = client.post("/api/auth/login", json={"username": "tmp_scoped", "password": "tmp-password-123"}).json()
    headers = {"Authorization": f"Bearer {token['access_token']}"}
    assert token["user"]["scope_label"] == "область не назначена"
    assert client.get("/api/channels", params={"at": AT}, headers=headers).json()["total"] == 0
    assert client.get("/api/dashboard/summary", params={"at": AT}, headers=headers).json()["channels_total"] == 0


def test_roles_add_up_and_scopes_union(client, auth, temp_user, session_factory, scope):
    with session_factory() as s:
        complexes = list(s.scalars(select(Object.id).where(Object.level == 2).order_by(Object.id)))
    two = [scope["complex_id"], next(c for c in complexes if c != scope["complex_id"])]
    r = client.put(f"/api/users/{temp_user}/access", headers=auth("admin"),
                   json={"roles": ["technician", "manager"], "scope_object_ids": two})
    assert r.status_code == 200, r.text
    assert r.json()["roles"] == ["manager", "technician"] and [n["id"] for n in r.json()["scope"]] == sorted(two)
    assert set(r.json()["permissions"]) == PERMISSIONS["technician"] | PERMISSIONS["manager"]
    token = client.post("/api/auth/login", json={"username": "tmp_scoped", "password": "tmp-password-123"}).json()
    headers = {"Authorization": f"Bearer {token['access_token']}"}
    tree = client.get("/api/objects", params={"at": AT}, headers=headers).json()["items"]
    assert sorted(c["id"] for c in tree[0]["children"]) == sorted(two)
    one = client.get("/api/channels", params={"at": AT, "limit": 500}, headers=h_tech(client)).json()["total"]
    both = client.get("/api/channels", params={"at": AT, "limit": 500}, headers=headers).json()["total"]
    assert both >= one
    # + диспетчер ОДС: видит всё
    client.put(f"/api/users/{temp_user}/access", headers=auth("admin"),
               json={"roles": ["technician", "dispatcher_ods"], "scope_object_ids": two})
    all_total = client.get("/api/channels", params={"at": AT}, headers=headers).json()["total"]
    assert all_total == client.get("/api/channels", params={"at": AT}, headers=auth("admin")).json()["total"]
    with session_factory() as s:
        last = s.scalars(select(AuditLog).where(AuditLog.action == "user_access_update")
                         .order_by(AuditLog.id.desc())).first()
        assert last.details["new"]["roles"] == ["dispatcher_ods", "technician"]
        assert last.details["old"]["roles"] == ["manager", "technician"]


def h_tech(client):
    token = client.post("/api/auth/login", json={"username": "technician", "password": "technician123"}).json()
    return {"Authorization": f"Bearer {token['access_token']}"}


def test_user_admin_validation(client, auth, temp_user, session_factory):
    admin = auth("admin")
    with session_factory() as s:
        sub = s.scalar(select(Object.id).where(Object.level == 3).limit(1))
        admin_id = s.scalar(select(User.id).where(User.username == "admin"))
    url = f"/api/users/{temp_user}/access"
    assert client.put(url, headers=admin, json={"roles": ["technician"], "scope_object_ids": []}).status_code == 422
    assert client.put(url, headers=admin, json={"roles": ["technician"], "scope_object_ids": [sub]}).status_code == 422
    assert client.put(url, headers=admin, json={"roles": ["technician"], "scope_object_ids": [999999]}).status_code == 422
    assert client.put(url, headers=admin, json={"roles": [], "scope_object_ids": []}).status_code == 422
    r = client.put(f"/api/users/{admin_id}/access", headers=admin, json={"roles": ["manager"], "scope_object_ids": [1]})
    assert r.status_code == 409
    assert client.put(url, headers=auth("dispatcher"), json={"roles": ["admin"]}).status_code == 403
    assert client.put("/api/users/999999/access", headers=admin, json={"roles": ["admin"]}).status_code == 404
    with session_factory() as s:
        u = s.get(User, temp_user)
        u.auth_source = "ldap"
        s.commit()
    assert client.put(url, headers=admin, json={"roles": ["admin"]}).status_code == 409
    with session_factory() as s:
        s.get(User, temp_user).auth_source = "local"
        s.commit()
    users = client.get("/api/users", headers=admin).json()["items"]
    assert {"ods", "technician", "dispatcher"} <= {u["username"] for u in users}


def test_full_access_constant_is_unrestricted(session_factory):
    with session_factory() as s:
        assert ObjectIndex.load(s, FULL_ACCESS).channel_objects is None
