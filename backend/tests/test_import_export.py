import io

from openpyxl import Workbook, load_workbook

KNOWN_CHANNEL = 334609


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_import_events_csv_applies_cleaning_rules(client, auth):
    # Заголовки как в файлах заказчика (с пробелами после извлечения из PDF), флаги f/t и false/true,
    # значение-дата, дубль в одну секунду, неизвестный канал
    csv_text = (
        "ид _ события;ид _ канала _ данных;дата;время;тревожное;значение _ датчика\n"
        f"101;{KNOWN_CHANNEL};2026-08-03;10:00:00;f;Норма\n"
        f"102;{KNOWN_CHANNEL};2026-08-03;10:00:01;t;Неисправен\n"
        f"103;{KNOWN_CHANNEL};2026-08-03;10:00:01;true;Неисправен\n"
        f"104;{KNOWN_CHANNEL};2026-08-03;10:00:02;false;01.01.1970 03:00:00\n"
        "105;1;2026-08-03;10:00:03;f;Норма\n"
        f"106;{KNOWN_CHANNEL};03.08.2026;10:00:04;false;0,015\n"
    )
    r = client.post("/api/import", data={"kind": "events"},
                    files={"file": ("journal.csv", csv_text.encode("utf-8-sig"), "text/csv")}, headers=auth("engineer"))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["rows_total"] == 6
    assert body["rows_imported"] == 3  # 101, 102, 106
    assert body["rows_skipped"] == 3  # дата вместо значения, неизвестный канал, дубль
    messages = " ".join(e["message"] for e in body["errors"])
    assert "Вместо показания записана дата" in messages
    assert "Канал 1 не найден" in messages

    # Повторная загрузка того же файла ничего не дублирует
    again = client.post("/api/import", data={"kind": "events"},
                        files={"file": ("journal.csv", csv_text.encode("utf-8"), "text/csv")}, headers=auth("engineer"))
    assert again.json()["rows_imported"] == 0

    hist = client.get(f"/api/channels/{KNOWN_CHANNEL}/history",
                      params={"at": "2026-08-03T12:00:00", "days": 1, "granularity": "hour"}, headers=auth()).json()
    assert any(f["ts"] == "2026-08-03T10:00:01" and f["kind"] == "Неисправен" for f in hist["faults"])


def test_import_channels_xlsx(client, auth):
    content = _xlsx([
        ["ид_канала_данных", "тип_инж_системы", "тип_датчика", "тег_инженерной_системы", "название_датчика", "ид_объект"],
        [999001, "Пожарная охрана", "Датчик дыма", "847-1.2.140.7.", "Дым ПК 1200+1", 25],
        [999002, "Вентиляция", "Вентилятор", "847-5.1.140.3", "Вентилятор ВШ-9 №2", 99999],
    ])
    r = client.post("/api/import", data={"kind": "channels"},
                    files={"file": ("channels.xlsx", content, "application/octet-stream")}, headers=auth("admin"))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["rows_imported"] == 1 and body["rows_skipped"] == 1
    card = client.get("/api/channels/999001", headers=auth()).json()
    assert card["tag_levels"] == ["847", "1", "2", "140", "7"]
    assert card["prediction"] is None


def test_import_predictions_from_ml(client, auth):
    csv_text = (
        "channel_id,at,horizon_h,prob,health,risk_level,top_factors,model_version\n"
        f'{KNOWN_CHANNEL},2026-09-01T00:00:00,24,0.91,9,critical,"[{{""feature"": ""fault_msgs_24h"", ""value"": 5, '
        f'""norm"": 0, ""phrase"": ""Сообщений о неисправности за сутки: 5 (обычно 0)""}}]",lgbm-test\n'
        f"{KNOWN_CHANNEL},2026-09-01T01:00:00,24,1.5,,,,lgbm-test\n"
    )
    r = client.post("/api/import", data={"kind": "predictions"},
                    files={"file": ("pred.csv", csv_text.encode(), "text/csv")}, headers=auth("admin"))
    assert r.status_code == 200, r.text
    assert r.json()["rows_imported"] == 1
    assert r.json()["errors"][0]["row"] == 3

    items = client.get("/api/predictions", params={"at": "2026-09-01T00:30:00"}, headers=auth()).json()["items"]
    assert items[0]["prob"] == 0.91
    assert items[0]["factors"] == ["Сообщений о неисправности за сутки: 5 (обычно 0)"]


def test_import_rejects_bad_files(client, auth):
    r = client.post("/api/import", data={"kind": "events"},
                    files={"file": ("x.txt", b"hello", "text/plain")}, headers=auth("admin"))
    assert r.status_code == 422
    assert r.json()["detail"] == "Поддерживаются только файлы .csv, .xlsx и .xml"

    r = client.post("/api/import", data={"kind": "events"},
                    files={"file": ("x.csv", b"foo,bar\n1,2\n", "text/csv")}, headers=auth("admin"))
    assert r.status_code == 422
    assert "нет обязательных колонок" in r.json()["detail"]

    r = client.post("/api/import", data={"kind": "nonsense"},
                    files={"file": ("x.csv", b"a\n1\n", "text/csv")}, headers=auth("admin"))
    assert r.status_code == 422


def test_export_journal_xlsx(client, auth):
    r = client.get("/api/export/journal", params={"at": "2026-08-03T00:00:00", "only_alerts": "false",
                                                  "date_from": "2026-08-02T00:00:00"}, headers=auth("manager"))
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert 'filename="journal_20260803_0000.xlsx"' in r.headers["content-disposition"]

    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:4] == ("Время прогноза", "ID прогноза", "ID датчика", "Датчик")
    assert rows[0][10:12] == ("Причины", "Причины (формулировки модели)")
    example = next(r for r in rows[1:] if r[10])
    assert example[10] != example[11]  # понятная формулировка и техническая — разные колонки
    # 60 датчиков × 5 срезов (02.08 00:00 … 03.08 00:00)
    assert len(rows) - 1 == 300

    journal = client.get("/api/journal", params={"at": "2026-08-03T00:00:00", "only_alerts": "false",
                                                 "date_from": "2026-08-02T00:00:00"}, headers=auth()).json()
    assert journal["total"] == 300


def test_ingest_events(client, auth):
    payload = {"events": [
        {"event_id": 1, "channel_id": KNOWN_CHANNEL, "ts": "2026-08-04T10:00:00", "is_alarm": "f", "value": "Норма"},
        {"channel_id": KNOWN_CHANNEL, "ts": "2026-08-04T10:00:00", "is_alarm": False, "value": "Норма"},
        {"channel_id": KNOWN_CHANNEL, "ts": "2026-08-04T10:01:00", "is_alarm": "t", "value": "01.01.1970 03:00:00"},
        {"channel_id": 1, "ts": "2026-08-04T10:02:00", "is_alarm": True, "value": 0.5},
    ]}
    r = client.post("/api/ingest/events", json=payload, headers=auth("admin"))
    assert r.status_code == 200, r.text
    assert r.json() == {
        "received": 4, "accepted": 1, "duplicates": 1, "skipped_unknown_channel": 1,
        "skipped_invalid": 1, "affected_channels": [KNOWN_CHANNEL],
    }
